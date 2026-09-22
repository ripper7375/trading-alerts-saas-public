"""
MCD3: Consolidated Trend and EDT Stochastic Evaluator (XAUUSD - M15 & M5)
DavinTrade Stack D - Engine 1.5A Module

Core Mandate:
1. Determine whether a strongly confirmed Consolidated Trend exists between M15 (macro) and M5 (micro):
   - Condition 1 (Trend Alignment): M15 trend direction == M5 trend direction (UPTREND, DOWNTREND, or SIDEWAYS, +/- 5.0° deadband).
   - Condition 2 (Historical Corridor Nesting): M5 EDT corridor [LOEDT, UOEDT] nested inside M15 EDT corridor
     for at least 75.0% of M5 EDT Time Horizon length (T_EDT, M5 bars).
   - Condition 3 (Current Bar Corridor Engulfment): M5 corridor completely engulfed within M15 corridor at Bar 0.
2. If all 3 conditions are satisfied -> Consolidated Trend confirmed.
   Calculate the Standard EDT Stochastic:
       EDT Stochastic = [(M15 SSA current - M15 LOEDT current) / (M15 UOEDT current - M15 LOEDT current)] * 100
   Intuitive Trader Scale:
       - 0.0% (LOEDT Floor): Price/SSA resting at lower support floor (Oversold / Deep Value Zone).
       - 50.0% (Corridor Midpoint): Price/SSA resting at equilibrium center.
       - 100.0% (UOEDT Ceiling): Price/SSA resting at upper resistance ceiling (Overbought / Climax Zone).
3. If any condition fails -> NON_CONSOLIDATED.
   EDT Stochastic is strictly UNAVAILABLE (null) to protect traders from erroneous execution.

4-Tier Pre-Flight Quality Gate Architecture:
- Tier 1: Candidate Isolation & Single Active Indicator Rule (M15: strictly 1 of 7 Centroids; M5: strictly 1 of 8 EDT indicators).
- Tier 2: Dual-Timeframe Time-Series Continuity & Non-Null Monotonicity.
- Tier 3: Dual Channel Boundary Sanity Gate (UOEDT > LOEDT on both timeframes for all evaluated bars).
- Tier 4: Dual Statistics Ingestion Verification (symbol=XAUUSD, containment_rate >= 50.0% on both horizons).

10 Discrete State Synthesis Matrix + System Error:
- Group 1: Consolidated Trend Confirmed (7 States across Bullish / Bearish / Sideways Stochastic Zones)
- Group 2: Non-Consolidated Trend (3 Diagnostic Failure States: Trend Conflict, Insufficient Nesting, Current Escape)
- Group 3: Data Pipeline Violation (MCD3_INVALID)
"""

import os
import sys
import json
import bisect
from datetime import datetime, timezone
from typing import Dict, Any, List, Optional, Tuple
import openpyxl

# Ensure UTF-8 output on Windows console
if sys.stdout and hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
if sys.stderr and hasattr(sys.stderr, "reconfigure"):
    try:
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass

# Candidate pools for M15 and M5
M15_CANDIDATE_INDICATORS = [
    "best_fit_a",
    "best_fit_b",
    "cherry_a",
    "cherry_b",
    "most_recent",
    "non_a",
    "non_b",
]

M5_CANDIDATE_INDICATORS = [
    "best_fit_a",
    "best_fit_b",
    "cherry_a",
    "cherry_b",
    "most_recent",
    "non_a",
    "non_b",
    "fractal",
]

# Configurable Calibration Constants
DEFAULT_SIDEWAYS_ANGLE_THRESHOLD = 5.0      # Deadband threshold in degrees (+/- 5.0°)
DEFAULT_MIN_CONTAINMENT_THRESHOLD = 50.0    # Corridor statistical integrity gate (%)
DEFAULT_MIN_NESTING_RATE_THRESHOLD = 75.0   # Minimum historical corridor nesting rate (%)
DEFAULT_M15_MIN_WINDOW_FLOOR = 96           # Minimum M15 bars floor: 96 bars (24 hours)
DEFAULT_M5_MIN_WINDOW_FLOOR = 48            # Minimum M5 bars floor: 48 bars (4 hours)
DEFAULT_STOCHASTIC_VALUE_THRESHOLD = 20.0   # Oversold / Value Dip Zone ceiling (%)
DEFAULT_STOCHASTIC_OVERBOUGHT_THRESHOLD = 80.0  # Overbought / Premium Zone floor (%)
DEFAULT_SYMBOL = "XAUUSD"


class MCD3ValidationError(Exception):
    """Custom exception for MCD3 dual-timeframe 4-tier validation failures."""
    pass


class MCD3ConsolidatedTrendEvaluator:
    """
    Evaluator for MCD3: Consolidated Trend and EDT Stochastic (XAUUSD - M15 & M5).
    """

    def __init__(
        self,
        excel_path: str,
        m15_target_indicator: Optional[str] = None,
        m5_target_indicator: Optional[str] = None,
        sideways_threshold: float = DEFAULT_SIDEWAYS_ANGLE_THRESHOLD,
        min_nesting_threshold: float = DEFAULT_MIN_NESTING_RATE_THRESHOLD,
        min_containment_threshold: float = DEFAULT_MIN_CONTAINMENT_THRESHOLD,
        m15_min_window_floor: int = DEFAULT_M15_MIN_WINDOW_FLOOR,
        m5_min_window_floor: int = DEFAULT_M5_MIN_WINDOW_FLOOR,
        stochastic_value_threshold: float = DEFAULT_STOCHASTIC_VALUE_THRESHOLD,
        stochastic_overbought_threshold: float = DEFAULT_STOCHASTIC_OVERBOUGHT_THRESHOLD,
        symbol: str = DEFAULT_SYMBOL,
    ):
        self.excel_path = os.path.abspath(excel_path)
        self.m15_target_indicator = m15_target_indicator.strip().lower() if m15_target_indicator else None
        self.m5_target_indicator = m5_target_indicator.strip().lower() if m5_target_indicator else None
        self.sideways_threshold = sideways_threshold
        self.min_nesting_threshold = min_nesting_threshold
        self.min_containment_threshold = min_containment_threshold
        self.m15_min_window_floor = m15_min_window_floor
        self.m5_min_window_floor = m5_min_window_floor
        self.stochastic_value_threshold = stochastic_value_threshold
        self.stochastic_overbought_threshold = stochastic_overbought_threshold
        self.symbol = symbol

        self.wb: Optional[openpyxl.Workbook] = None
        self.m15_sheet = None
        self.m5_sheet = None
        self.stat_sheet = None

        self.validation_errors: List[str] = []
        self.validation_warnings: List[str] = []
        self.validation_checks: Dict[str, Any] = {}

    def load_workbook(self) -> None:
        """Load Excel workbook in read-only mode for high performance."""
        if not os.path.exists(self.excel_path):
            raise FileNotFoundError(f"Excel workbook not found at: {self.excel_path}")

        try:
            self.wb = openpyxl.load_workbook(self.excel_path, data_only=True, read_only=True)
        except Exception as e:
            raise MCD3ValidationError(f"Failed to open Excel workbook: {str(e)}")

        sheet_names = self.wb.sheetnames
        m15_sheet_name = "market_data_v6_M15"
        m5_sheet_name = "market_data_v6_M5"

        if m15_sheet_name not in sheet_names:
            raise MCD3ValidationError(
                f"Required sheet '{m15_sheet_name}' not found. Available sheets: {sheet_names}"
            )
        if m5_sheet_name not in sheet_names:
            raise MCD3ValidationError(
                f"Required sheet '{m5_sheet_name}' not found. Available sheets: {sheet_names}"
            )
        if "indicator_statistics" not in sheet_names:
            raise MCD3ValidationError(
                f"Required sheet 'indicator_statistics' not found. Available sheets: {sheet_names}"
            )

        self.m15_sheet = self.wb[m15_sheet_name]
        self.m5_sheet = self.wb[m5_sheet_name]
        self.stat_sheet = self.wb["indicator_statistics"]

    def close(self) -> None:
        """Close Excel workbook to release Windows file handles."""
        if self.wb is not None:
            try:
                self.wb.close()
            except Exception:
                pass
            self.wb = None

    def _get_sheet_headers_and_rows(self, sheet) -> Tuple[Dict[str, int], List[List[Any]]]:
        """Extract header column map and rows from an openpyxl sheet."""
        rows_iter = sheet.iter_rows(values_only=True)
        try:
            header_row = next(rows_iter)
        except StopIteration:
            return {}, []

        headers = {str(name).strip(): col_idx for col_idx, name in enumerate(header_row) if name is not None}
        data_rows = [list(r) for r in rows_iter if any(v is not None for v in r)]
        return headers, data_rows

    def validate_tier1_m15(
        self, headers: Dict[str, int], rows: List[List[Any]]
    ) -> Tuple[str, int]:
        """
        Tier 1 Validation for M15: Candidate Isolation & Single Active Indicator Rule.
        Scans strictly the 7 Centroid Candidates on M15.
        Enforces exactly 1 active indicator in production mode.
        """
        total_bars = len(rows)
        self.validation_checks["m15_total_bars"] = total_bars

        if total_bars < self.m15_min_window_floor:
            err = (
                f"Tier 1 Validation Failed on M15: Insufficient bar data in market_data_v6_M15. "
                f"Found {total_bars} bars, but minimum floor requires {self.m15_min_window_floor} bars."
            )
            self.validation_errors.append(err)
            raise MCD3ValidationError(err)

        detected_active: List[str] = []
        indicator_coverage: Dict[str, int] = {}
        indicator_last_row_indices: Dict[str, int] = {}

        for ind in M15_CANDIDATE_INDICATORS:
            ssa_col = f"{ind}_ssa"
            uoedt_col = f"{ind}_uoedt"
            loedt_col = f"{ind}_loedt"

            if ssa_col not in headers or uoedt_col not in headers or loedt_col not in headers:
                indicator_coverage[ind] = 0
                continue

            ssa_idx = headers[ssa_col]
            uoedt_idx = headers[uoedt_col]
            loedt_idx = headers[loedt_col]

            populated_indices = [
                i
                for i, r in enumerate(rows)
                if r[ssa_idx] is not None
                and r[uoedt_idx] is not None
                and r[loedt_idx] is not None
                and isinstance(r[ssa_idx], (int, float))
                and isinstance(r[uoedt_idx], (int, float))
                and isinstance(r[loedt_idx], (int, float))
                and r[ssa_idx] > 0
            ]

            count = len(populated_indices)
            indicator_coverage[ind] = count
            if count >= self.m15_min_window_floor:
                detected_active.append(ind)
                indicator_last_row_indices[ind] = populated_indices[-1]
            elif count > 0:
                self.validation_warnings.append(
                    f"Insufficient bar coverage for M15 '{ind}': found {count} bars, minimum floor is {self.m15_min_window_floor}."
                )

        self.validation_checks["m15_candidate_coverage"] = indicator_coverage
        self.validation_checks["m15_active_detected"] = detected_active

        # If a target indicator was explicitly specified (for test isolation)
        if self.m15_target_indicator is not None:
            if self.m15_target_indicator not in M15_CANDIDATE_INDICATORS:
                err = (
                    f"Tier 1 Validation Failed: Specified m15_target_indicator '{self.m15_target_indicator}' "
                    f"is not in permitted M15 candidates: {M15_CANDIDATE_INDICATORS}"
                )
                self.validation_errors.append(err)
                raise MCD3ValidationError(err)

            if self.m15_target_indicator not in detected_active:
                err = (
                    f"Tier 1 Validation Failed: Specified m15_target_indicator '{self.m15_target_indicator}' "
                    f"has insufficient data ({indicator_coverage.get(self.m15_target_indicator, 0)} bars)."
                )
                self.validation_errors.append(err)
                raise MCD3ValidationError(err)

            return self.m15_target_indicator, indicator_last_row_indices[self.m15_target_indicator]

        # Production Mode: Exactly 1 active indicator
        if len(detected_active) == 0:
            err = (
                "Tier 1 Validation Failed on M15: No active Centroid indicator detected with valid data "
                f"covering at least {self.m15_min_window_floor} bars on M15."
            )
            self.validation_errors.append(err)
            raise MCD3ValidationError(err)

        if len(detected_active) > 1:
            err = (
                f"Tier 1 Validation Failed on M15: Multiple active indicators detected on M15: {detected_active}. "
                "Production safety mandate requires strictly 1 active indicator."
            )
            self.validation_errors.append(err)
            raise MCD3ValidationError(err)

        active_ind = detected_active[0]
        return active_ind, indicator_last_row_indices[active_ind]

    def validate_tier1_m5(
        self, headers: Dict[str, int], rows: List[List[Any]]
    ) -> Tuple[str, int]:
        """
        Tier 1 Validation for M5: Candidate Isolation & Single Active Indicator Rule.
        Scans strictly the 8 Candidate EDT Indicators on M5.
        Enforces exactly 1 active indicator in production mode.
        """
        total_bars = len(rows)
        self.validation_checks["m5_total_bars"] = total_bars

        if total_bars < self.m5_min_window_floor:
            err = (
                f"Tier 1 Validation Failed on M5: Insufficient bar data in market_data_v6_M5. "
                f"Found {total_bars} bars, but minimum floor requires {self.m5_min_window_floor} bars."
            )
            self.validation_errors.append(err)
            raise MCD3ValidationError(err)

        detected_active: List[str] = []
        indicator_coverage: Dict[str, int] = {}
        indicator_last_row_indices: Dict[str, int] = {}

        for ind in M5_CANDIDATE_INDICATORS:
            if ind == "fractal":
                fl_col = "fractal_best_fl"
                uoedt_col = "fractal_uoedt"
                loedt_col = "fractal_loedt"
                close_col = "close"

                if fl_col not in headers or uoedt_col not in headers or loedt_col not in headers:
                    indicator_coverage[ind] = 0
                    continue

                fl_idx = headers[fl_col]
                uoedt_idx = headers[uoedt_col]
                loedt_idx = headers[loedt_col]
                close_idx = headers.get(close_col, -1)

                populated_indices = [
                    i
                    for i, r in enumerate(rows)
                    if r[fl_idx] is not None
                    and r[uoedt_idx] is not None
                    and r[loedt_idx] is not None
                    and (close_idx == -1 or r[close_idx] is not None)
                    and isinstance(r[fl_idx], (int, float))
                    and isinstance(r[uoedt_idx], (int, float))
                    and isinstance(r[loedt_idx], (int, float))
                    and r[uoedt_idx] > 0
                ]
            else:
                ssa_col = f"{ind}_ssa"
                uoedt_col = f"{ind}_uoedt"
                loedt_col = f"{ind}_loedt"

                if ssa_col not in headers or uoedt_col not in headers or loedt_col not in headers:
                    indicator_coverage[ind] = 0
                    continue

                ssa_idx = headers[ssa_col]
                uoedt_idx = headers[uoedt_col]
                loedt_idx = headers[loedt_col]

                populated_indices = [
                    i
                    for i, r in enumerate(rows)
                    if r[ssa_idx] is not None
                    and r[uoedt_idx] is not None
                    and r[loedt_idx] is not None
                    and isinstance(r[ssa_idx], (int, float))
                    and isinstance(r[uoedt_idx], (int, float))
                    and isinstance(r[loedt_idx], (int, float))
                    and r[ssa_idx] > 0
                ]

            count = len(populated_indices)
            indicator_coverage[ind] = count
            if count >= self.m5_min_window_floor:
                detected_active.append(ind)
                indicator_last_row_indices[ind] = populated_indices[-1]
            elif count > 0:
                self.validation_warnings.append(
                    f"Insufficient bar coverage for M5 '{ind}': found {count} bars, minimum floor is {self.m5_min_window_floor}."
                )

        self.validation_checks["m5_candidate_coverage"] = indicator_coverage
        self.validation_checks["m5_active_detected"] = detected_active

        # If a target indicator was explicitly specified (for test isolation)
        if self.m5_target_indicator is not None:
            if self.m5_target_indicator not in M5_CANDIDATE_INDICATORS:
                err = (
                    f"Tier 1 Validation Failed: Specified m5_target_indicator '{self.m5_target_indicator}' "
                    f"is not in permitted M5 candidates: {M5_CANDIDATE_INDICATORS}"
                )
                self.validation_errors.append(err)
                raise MCD3ValidationError(err)

            if self.m5_target_indicator not in detected_active:
                err = (
                    f"Tier 1 Validation Failed: Specified m5_target_indicator '{self.m5_target_indicator}' "
                    f"has insufficient data ({indicator_coverage.get(self.m5_target_indicator, 0)} bars)."
                )
                self.validation_errors.append(err)
                raise MCD3ValidationError(err)

            return self.m5_target_indicator, indicator_last_row_indices[self.m5_target_indicator]

        # Production Mode: Exactly 1 active indicator
        if len(detected_active) == 0:
            err = (
                "Tier 1 Validation Failed on M5: No active EDT indicator detected with valid data "
                f"covering at least {self.m5_min_window_floor} bars on M5."
            )
            self.validation_errors.append(err)
            raise MCD3ValidationError(err)

        if len(detected_active) > 1:
            err = (
                f"Tier 1 Validation Failed on M5: Multiple active indicators detected on M5: {detected_active}. "
                "Production safety mandate requires strictly 1 active indicator."
            )
            self.validation_errors.append(err)
            raise MCD3ValidationError(err)

        active_ind = detected_active[0]
        return active_ind, indicator_last_row_indices[active_ind]

    def validate_tier4_statistics(
        self,
        stat_headers: Dict[str, int],
        stat_rows: List[List[Any]],
        m15_active: str,
        m5_active: str,
    ) -> Tuple[Dict[str, Any], Dict[str, Any]]:
        """
        Tier 4 Validation: Dual Statistics Ingestion & Channel Quality Integrity.
        Queries indicator_statistics for symbol='XAUUSD' on both M15 and M5.
        Asserts containment_rate >= 50.0% for both active indicators.
        """
        required_cols = ["symbol", "timeframe", "source", "regression_angle", "containment_rate"]
        for col in required_cols:
            if col not in stat_headers:
                err = f"Tier 4 Validation Failed: Required column '{col}' missing from indicator_statistics."
                self.validation_errors.append(err)
                raise MCD3ValidationError(err)

        sym_idx = stat_headers["symbol"]
        tf_idx = stat_headers["timeframe"]
        src_idx = stat_headers["source"]
        angle_idx = stat_headers["regression_angle"]
        cr_idx = stat_headers["containment_rate"]
        n_idx = stat_headers.get("containment_n", -1)

        # 1. Ingest M15 stats
        m15_stat_row = None
        for r in stat_rows:
            if (
                r[sym_idx] == self.symbol
                and str(r[tf_idx]).strip().upper() == "M15"
                and str(r[src_idx]).strip().lower() == m15_active
            ):
                m15_stat_row = r
                break

        if m15_stat_row is None:
            err = (
                f"Tier 4 Validation Failed on M15: No statistics record found in indicator_statistics "
                f"for symbol='{self.symbol}', timeframe='M15', source='{m15_active}'."
            )
            self.validation_errors.append(err)
            raise MCD3ValidationError(err)

        m15_angle = m15_stat_row[angle_idx]
        m15_cr = m15_stat_row[cr_idx]
        m15_n = m15_stat_row[n_idx] if n_idx != -1 else None

        if m15_angle is None or not isinstance(m15_angle, (int, float)):
            err = f"Tier 4 Validation Failed: Invalid regression_angle for M15 '{m15_active}': {m15_angle}"
            self.validation_errors.append(err)
            raise MCD3ValidationError(err)

        if m15_cr is None or not isinstance(m15_cr, (int, float)):
            err = f"Tier 4 Validation Failed: Invalid containment_rate for M15 '{m15_active}': {m15_cr}"
            self.validation_errors.append(err)
            raise MCD3ValidationError(err)

        if m15_cr < self.min_containment_threshold:
            err = (
                f"Tier 4 Validation Failed on M15: Containment rate for '{m15_active}' is {m15_cr:.2f}%, "
                f"which is below the minimum threshold of {self.min_containment_threshold}%."
            )
            self.validation_errors.append(err)
            raise MCD3ValidationError(err)

        # 2. Ingest M5 stats (fractal maps to 'fractal_edt' in stats)
        m5_expected_source = "fractal_edt" if m5_active == "fractal" else m5_active
        m5_stat_row = None
        for r in stat_rows:
            if (
                r[sym_idx] == self.symbol
                and str(r[tf_idx]).strip().upper() == "M5"
                and str(r[src_idx]).strip().lower() == m5_expected_source
            ):
                m5_stat_row = r
                break

        if m5_stat_row is None:
            err = (
                f"Tier 4 Validation Failed on M5: No statistics record found in indicator_statistics "
                f"for symbol='{self.symbol}', timeframe='M5', source='{m5_expected_source}'."
            )
            self.validation_errors.append(err)
            raise MCD3ValidationError(err)

        m5_angle = m5_stat_row[angle_idx]
        m5_cr = m5_stat_row[cr_idx]
        m5_n = m5_stat_row[n_idx] if n_idx != -1 else None

        if m5_angle is None or not isinstance(m5_angle, (int, float)):
            err = f"Tier 4 Validation Failed: Invalid regression_angle for M5 '{m5_active}': {m5_angle}"
            self.validation_errors.append(err)
            raise MCD3ValidationError(err)

        if m5_cr is None or not isinstance(m5_cr, (int, float)):
            err = f"Tier 4 Validation Failed: Invalid containment_rate for M5 '{m5_active}': {m5_cr}"
            self.validation_errors.append(err)
            raise MCD3ValidationError(err)

        if m5_cr < self.min_containment_threshold:
            err = (
                f"Tier 4 Validation Failed on M5: Containment rate for '{m5_active}' is {m5_cr:.2f}%, "
                f"which is below the minimum threshold of {self.min_containment_threshold}%."
            )
            self.validation_errors.append(err)
            raise MCD3ValidationError(err)

        m15_metrics = {
            "source": m15_active,
            "regression_angle": float(m15_angle),
            "containment_rate": float(m15_cr),
            "containment_n": int(m15_n) if m15_n is not None else 1808,
        }

        m5_metrics = {
            "source": m5_active,
            "regression_angle": float(m5_angle),
            "containment_rate": float(m5_cr),
            "containment_n": int(m5_n) if m5_n is not None else 755,
        }

        return m15_metrics, m5_metrics

    def validate_tier2_and_tier3_continuity_and_sanity(
        self,
        m15_headers: Dict[str, int],
        m15_rows: List[List[Any]],
        m15_active: str,
        m5_headers: Dict[str, int],
        m5_rows: List[List[Any]],
        m5_active: str,
    ) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
        """
        Tier 2 Validation: Dual Time-Series Continuity & Non-Null Monotonicity.
        Tier 3 Validation: Dual Channel Boundary Sanity Gate (UOEDT > LOEDT on both horizons).
        Returns filtered, validated candle objects for M15 and M5.
        """
        # --- M15 Validation ---
        m15_ts_idx = m15_headers.get("timestamp", -1)
        m15_close_idx = m15_headers.get("close", -1)
        m15_ssa_idx = m15_headers.get(f"{m15_active}_ssa", -1)
        m15_uoedt_idx = m15_headers.get(f"{m15_active}_uoedt", -1)
        m15_loedt_idx = m15_headers.get(f"{m15_active}_loedt", -1)

        if min(m15_ts_idx, m15_close_idx, m15_ssa_idx, m15_uoedt_idx, m15_loedt_idx) < 0:
            err = "Tier 2 Validation Failed on M15: Missing required column indices in M15 header."
            self.validation_errors.append(err)
            raise MCD3ValidationError(err)

        m15_bars: List[Dict[str, Any]] = []
        prev_ts = None
        for i, r in enumerate(m15_rows):
            ts = r[m15_ts_idx]
            close = r[m15_close_idx]
            ssa = r[m15_ssa_idx]
            uoedt = r[m15_uoedt_idx]
            loedt = r[m15_loedt_idx]

            # Skip unpopulated initial bars if any
            if ssa is None or uoedt is None or loedt is None:
                continue

            if not isinstance(ts, (int, float)) or not isinstance(close, (int, float)):
                continue

            # Tier 2 Monotonicity
            if prev_ts is not None and ts <= prev_ts:
                err = f"Tier 2 Validation Failed on M15: Timestamp non-monotonic at bar {i}: {ts} <= {prev_ts}."
                self.validation_errors.append(err)
                raise MCD3ValidationError(err)
            prev_ts = ts

            # Tier 3 Channel Sanity
            if uoedt <= loedt:
                err = f"Tier 3 Validation Failed on M15: Channel inverted or zero-width at bar {i}: UOEDT={uoedt} <= LOEDT={loedt}."
                self.validation_errors.append(err)
                raise MCD3ValidationError(err)

            m15_bars.append({
                "index": i,
                "timestamp": int(ts),
                "close": float(close),
                "ssa": float(ssa),
                "uoedt": float(uoedt),
                "loedt": float(loedt),
            })

        if len(m15_bars) < self.m15_min_window_floor:
            err = f"Tier 2 Validation Failed on M15: Valid bar count {len(m15_bars)} < {self.m15_min_window_floor}."
            self.validation_errors.append(err)
            raise MCD3ValidationError(err)

        # --- M5 Validation ---
        m5_ts_idx = m5_headers.get("timestamp", -1)
        m5_close_idx = m5_headers.get("close", -1)
        if m5_active == "fractal":
            m5_val_idx = m5_close_idx  # Fractal uses close
            m5_uoedt_idx = m5_headers.get("fractal_uoedt", -1)
            m5_loedt_idx = m5_headers.get("fractal_loedt", -1)
        else:
            m5_val_idx = m5_headers.get(f"{m5_active}_ssa", -1)
            m5_uoedt_idx = m5_headers.get(f"{m5_active}_uoedt", -1)
            m5_loedt_idx = m5_headers.get(f"{m5_active}_loedt", -1)

        if min(m5_ts_idx, m5_close_idx, m5_val_idx, m5_uoedt_idx, m5_loedt_idx) < 0:
            err = "Tier 2 Validation Failed on M5: Missing required column indices in M5 header."
            self.validation_errors.append(err)
            raise MCD3ValidationError(err)

        m5_bars: List[Dict[str, Any]] = []
        prev_ts = None
        for i, r in enumerate(m5_rows):
            ts = r[m5_ts_idx]
            close = r[m5_close_idx]
            val = r[m5_val_idx]
            uoedt = r[m5_uoedt_idx]
            loedt = r[m5_loedt_idx]

            if val is None or uoedt is None or loedt is None:
                continue

            if not isinstance(ts, (int, float)) or not isinstance(close, (int, float)):
                continue

            # Tier 2 Monotonicity
            if prev_ts is not None and ts <= prev_ts:
                err = f"Tier 2 Validation Failed on M5: Timestamp non-monotonic at bar {i}: {ts} <= {prev_ts}."
                self.validation_errors.append(err)
                raise MCD3ValidationError(err)
            prev_ts = ts

            # Tier 3 Channel Sanity
            if uoedt <= loedt:
                err = f"Tier 3 Validation Failed on M5: Channel inverted or zero-width at bar {i}: UOEDT={uoedt} <= LOEDT={loedt}."
                self.validation_errors.append(err)
                raise MCD3ValidationError(err)

            m5_bars.append({
                "index": i,
                "timestamp": int(ts),
                "close": float(close),
                "value": float(val),
                "uoedt": float(uoedt),
                "loedt": float(loedt),
            })

        if len(m5_bars) < self.m5_min_window_floor:
            err = f"Tier 2 Validation Failed on M5: Valid bar count {len(m5_bars)} < {self.m5_min_window_floor}."
            self.validation_errors.append(err)
            raise MCD3ValidationError(err)

        # Check chronological overlap between M15 and M5
        if m5_bars[-1]["timestamp"] < m15_bars[0]["timestamp"] or m15_bars[-1]["timestamp"] < m5_bars[0]["timestamp"]:
            err = "Tier 2 Validation Failed: No chronological time overlap between M15 and M5 datasets."
            self.validation_errors.append(err)
            raise MCD3ValidationError(err)

        return m15_bars, m5_bars

    def classify_trend(self, angle: float) -> str:
        """Classify trend direction from linear regression angle using deadband threshold."""
        if angle > self.sideways_threshold:
            return "UPTREND"
        elif angle < -self.sideways_threshold:
            return "DOWNTREND"
        else:
            return "SIDEWAYS"

    def evaluate(self) -> Dict[str, Any]:
        """
        Execute full MCD3 evaluation pipeline across both M15 and M5 horizons.
        Returns standard DavinTrade Stack D JSONB payload.
        """
        try:
            self.load_workbook()
            m15_headers, m15_rows = self._get_sheet_headers_and_rows(self.m15_sheet)
            m5_headers, m5_rows = self._get_sheet_headers_and_rows(self.m5_sheet)
            stat_headers, stat_rows = self._get_sheet_headers_and_rows(self.stat_sheet)
            self.close()

            # Tier 1 Validation
            m15_active, _ = self.validate_tier1_m15(m15_headers, m15_rows)
            m5_active, _ = self.validate_tier1_m5(m5_headers, m5_rows)

            # Tier 4 Validation (Statistics)
            m15_stats, m5_stats = self.validate_tier4_statistics(
                stat_headers, stat_rows, m15_active, m5_active
            )

            # Tier 2 & Tier 3 Validation (Continuity & Sanity)
            m15_bars, m5_bars = self.validate_tier2_and_tier3_continuity_and_sanity(
                m15_headers, m15_rows, m15_active, m5_headers, m5_rows, m5_active
            )

        except Exception as e:
            self.close()
            # Pre-flight validation failure -> emit INVALID state safely
            return self._build_invalid_payload(str(e))

        # --- MCD3 Core Calculations ---
        # 1. Condition 1: Trend Alignment
        m15_angle = m15_stats["regression_angle"]
        m5_angle = m5_stats["regression_angle"]
        m15_trend = self.classify_trend(m15_angle)
        m5_trend = self.classify_trend(m5_angle)

        condition_1_trend_aligned = (m15_trend == m5_trend)

        # 2. Condition 2: Historical Corridor Nesting (T_EDT, M5 bars)
        # Build sorted list of M15 timestamps for backward as-of lookup
        m15_timestamps = [b["timestamp"] for b in m15_bars]
        m5_horizon_n = m5_stats["containment_n"]

        # Evaluate over the last min(m5_horizon_n, len(m5_bars)) bars
        eval_span = min(m5_horizon_n, len(m5_bars))
        eval_m5_bars = m5_bars[-eval_span:]

        contained_bars = 0
        total_eval_bars = len(eval_m5_bars)

        for m5_b in eval_m5_bars:
            t_m5 = m5_b["timestamp"]
            idx = bisect.bisect_right(m15_timestamps, t_m5) - 1
            if idx >= 0:
                m15_b = m15_bars[idx]
                # Check if M5 corridor is nested inside M15 corridor
                if m5_b["loedt"] >= m15_b["loedt"] and m5_b["uoedt"] <= m15_b["uoedt"]:
                    contained_bars += 1

        nesting_rate_pct = (contained_bars / total_eval_bars * 100.0) if total_eval_bars > 0 else 0.0
        condition_2_historical_nesting_passed = (nesting_rate_pct >= self.min_nesting_threshold)

        # 3. Condition 3: Current Bar Complete Corridor Engulfment (Bar 0)
        latest_m5 = m5_bars[-1]
        latest_m15 = m15_bars[-1]

        condition_3_current_bar_engulfed = (
            latest_m5["loedt"] >= latest_m15["loedt"]
            and latest_m5["uoedt"] <= latest_m15["uoedt"]
        )

        # Consolidated Trend Flag
        is_consolidated_trend = (
            condition_1_trend_aligned
            and condition_2_historical_nesting_passed
            and condition_3_current_bar_engulfed
        )

        # Consolidated Trend State Name
        if is_consolidated_trend:
            if m15_trend == "UPTREND":
                consolidated_trend_state = "BULLISH_CONSOLIDATED"
            elif m15_trend == "DOWNTREND":
                consolidated_trend_state = "BEARISH_CONSOLIDATED"
            else:
                consolidated_trend_state = "SIDEWAYS_CONSOLIDATED"
        else:
            consolidated_trend_state = "NON_CONSOLIDATED"

        # 4. Standard EDT Stochastic (Conditional Execution)
        edt_stochastic: Optional[float] = None
        stochastic_zone = "UNAVAILABLE"
        regime_status = "UNCERTAIN"
        discrete_state_code = "MCD3_UNKNOWN"
        tactical_bias = "NEUTRAL_STAND_ASIDE"

        if is_consolidated_trend:
            # Standard EDT Stochastic: [(M15 SSA curr - M15 LOEDT curr) / (M15 UOEDT curr - M15 LOEDT curr)] * 100
            m15_corridor_width = latest_m15["uoedt"] - latest_m15["loedt"]
            if m15_corridor_width > 0:
                raw_stoch = ((latest_m15["ssa"] - latest_m15["loedt"]) / m15_corridor_width) * 100.0
                edt_stochastic = round(raw_stoch, 2)
            else:
                edt_stochastic = 50.0

            # State Synthesis across the 7 Consolidated States
            if consolidated_trend_state == "BULLISH_CONSOLIDATED":
                if edt_stochastic <= self.stochastic_value_threshold:
                    stochastic_zone = "VALUE_ZONE"
                    regime_status = "BULLISH_CONSOLIDATED_VALUE_ZONE"
                    discrete_state_code = "MCD3_BULL_VALUE"
                    tactical_bias = "HIGH_CONVICTION_BUY_DIP"
                elif edt_stochastic >= self.stochastic_overbought_threshold:
                    stochastic_zone = "OVERBOUGHT_ZONE"
                    regime_status = "BULLISH_CONSOLIDATED_OVERBOUGHT"
                    discrete_state_code = "MCD3_BULL_TOP"
                    tactical_bias = "CAUTION_TAKE_PROFIT_BUY"
                else:
                    stochastic_zone = "EQUILIBRIUM"
                    regime_status = "BULLISH_CONSOLIDATED_EQUILIBRIUM"
                    discrete_state_code = "MCD3_BULL_MID"
                    tactical_bias = "HOLD_BULLISH_TREND_RUNNER"

            elif consolidated_trend_state == "BEARISH_CONSOLIDATED":
                if edt_stochastic >= self.stochastic_overbought_threshold:
                    stochastic_zone = "PREMIUM_ZONE"
                    regime_status = "BEARISH_CONSOLIDATED_PREMIUM_ZONE"
                    discrete_state_code = "MCD3_BEAR_PREMIUM"
                    tactical_bias = "HIGH_CONVICTION_SELL_RALLY"
                elif edt_stochastic <= self.stochastic_value_threshold:
                    stochastic_zone = "OVERSOLD_ZONE"
                    regime_status = "BEARISH_CONSOLIDATED_OVERSOLD"
                    discrete_state_code = "MCD3_BEAR_BOTTOM"
                    tactical_bias = "CAUTION_TAKE_PROFIT_SELL"
                else:
                    stochastic_zone = "EQUILIBRIUM"
                    regime_status = "BEARISH_CONSOLIDATED_EQUILIBRIUM"
                    discrete_state_code = "MCD3_BEAR_MID"
                    tactical_bias = "HOLD_BEARISH_TREND_RUNNER"

            else:  # SIDEWAYS_CONSOLIDATED
                stochastic_zone = "IN_CORRIDOR"
                regime_status = "SIDEWAYS_CONSOLIDATED_EQUILIBRIUM"
                discrete_state_code = "MCD3_SIDEWAYS_EQUILIBRIUM"
                tactical_bias = "RANGE_BOUND_MEAN_REVERSION"

        else:
            # Non-Consolidated: Stochastic strictly null; isolate failure cause
            edt_stochastic = None
            stochastic_zone = "UNAVAILABLE"
            tactical_bias = "NEUTRAL_STAND_ASIDE"

            if not condition_1_trend_aligned:
                regime_status = "TREND_MISALIGNMENT"
                discrete_state_code = "MCD3_NON_CONSOLIDATED_TREND_CONFLICT"
            elif not condition_2_historical_nesting_passed:
                regime_status = "INSUFFICIENT_CORRIDOR_NESTING"
                discrete_state_code = "MCD3_NON_CONSOLIDATED_OVERFLOW"
            else:  # condition 3 failed
                regime_status = "CURRENT_CORRIDOR_ESCAPE"
                discrete_state_code = "MCD3_NON_CONSOLIDATED_ESCAPE"

        # 5. Zero-Hallucination Canonical English Commentary
        commentary = self._generate_canonical_commentary(
            discrete_state_code=discrete_state_code,
            m15_trend=m15_trend,
            m15_angle=m15_angle,
            m5_trend=m5_trend,
            m5_angle=m5_angle,
            nesting_pct=nesting_rate_pct,
            contained_bars=contained_bars,
            total_bars=total_eval_bars,
            stoch=edt_stochastic,
            latest_m5=latest_m5,
            latest_m15=latest_m15,
        )

        # Convert latest bar timestamp to ISO UTC string
        try:
            iso_timestamp = datetime.fromtimestamp(latest_m5["timestamp"], tz=timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
        except Exception:
            iso_timestamp = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")

        # Build production JSONB payload
        output_payload: Dict[str, Any] = {
            "symbol": self.symbol,
            "timeframe": "M15_M5",
            "timestamp": iso_timestamp,
            "evaluator_version": "1.0.0",
            "module": "MCD3_CONSOLIDATED_TREND_AND_EDT_STOCHASTIC",
            "validation": {
                "status": "PASS",
                "tier1_active_indicators": {
                    "m15": [m15_active],
                    "m5": [m5_active],
                },
                "tier2_continuity": "PASS",
                "tier3_sanity": "PASS",
                "tier4_statistics": "PASS",
                "warnings": self.validation_warnings,
            },
            "m15_metrics": {
                "active_indicator": m15_active,
                "trend_state": m15_trend,
                "regression_angle": round(m15_angle, 2),
                "current_close": round(latest_m15["close"], 2),
                "current_ssa": round(latest_m15["ssa"], 2),
                "current_uoedt": round(latest_m15["uoedt"], 2),
                "current_loedt": round(latest_m15["loedt"], 2),
                "containment_rate": round(m15_stats["containment_rate"], 2),
                "edt_horizon_n": m15_stats["containment_n"],
            },
            "m5_metrics": {
                "active_indicator": m5_active,
                "trend_state": m5_trend,
                "regression_angle": round(m5_angle, 2),
                "current_close": round(latest_m5["close"], 2),
                "current_ssa_or_close": round(latest_m5["value"], 2),
                "current_uoedt": round(latest_m5["uoedt"], 2),
                "current_loedt": round(latest_m5["loedt"], 2),
                "containment_rate": round(m5_stats["containment_rate"], 2),
                "edt_horizon_n": m5_stats["containment_n"],
            },
            "consolidated_trend_conditions": {
                "condition_1_trend_aligned": condition_1_trend_aligned,
                "condition_2_historical_nesting_passed": condition_2_historical_nesting_passed,
                "condition_2_nesting_rate_pct": round(nesting_rate_pct, 2),
                "condition_2_contained_bars": contained_bars,
                "condition_2_total_bars": total_eval_bars,
                "condition_3_current_bar_engulfed": condition_3_current_bar_engulfed,
            },
            "evaluation": {
                "is_consolidated_trend": is_consolidated_trend,
                "consolidated_trend_state": consolidated_trend_state,
                "edt_stochastic": edt_stochastic,
                "stochastic_zone": stochastic_zone,
                "regime_status": regime_status,
                "discrete_state_code": discrete_state_code,
                "tactical_bias": tactical_bias,
            },
            "commentary": commentary,
        }

        return output_payload

    def _generate_canonical_commentary(
        self,
        discrete_state_code: str,
        m15_trend: str,
        m15_angle: float,
        m5_trend: str,
        m5_angle: float,
        nesting_pct: float,
        contained_bars: int,
        total_bars: int,
        stoch: Optional[float],
        latest_m5: Dict[str, Any],
        latest_m15: Dict[str, Any],
    ) -> str:
        """Format zero-hallucination canonical English commentary from strict templates."""
        m15_sign = "+" if m15_angle > 0 else ""
        m5_sign = "+" if m5_angle > 0 else ""

        if discrete_state_code == "MCD3_NON_CONSOLIDATED_TREND_CONFLICT":
            return (
                f"XAUUSD multi-timeframe structure is NON_CONSOLIDATED due to trend conflict: "
                f"M15 slope is {m15_trend} ({m15_sign}{m15_angle:.2f}°) while M5 slope is {m5_trend} "
                f"({m5_sign}{m5_angle:.2f}°). Because the 3 strict conditions are not satisfied, "
                f"a Consolidated Trend does not exist and EDT Stochastic is UNAVAILABLE."
            )
        elif discrete_state_code == "MCD3_NON_CONSOLIDATED_OVERFLOW":
            return (
                f"XAUUSD multi-timeframe trends are aligned ({m15_trend}), but historical M5 corridor "
                f"nesting within M15 corridor is {nesting_pct:.1f}%, which fails the strict 75.0% threshold requirement "
                f"({contained_bars}/{total_bars} bars). A Consolidated Trend cannot be confirmed and EDT Stochastic is UNAVAILABLE."
            )
        elif discrete_state_code == "MCD3_NON_CONSOLIDATED_ESCAPE":
            return (
                f"XAUUSD multi-timeframe trends are aligned ({m15_trend}) with {nesting_pct:.1f}% historical nesting, "
                f"but the M5 corridor on the current bar extends outside the M15 corridor boundaries "
                f"(M5: [{latest_m5['loedt']:.2f}, {latest_m5['uoedt']:.2f}] vs M15: [{latest_m15['loedt']:.2f}, {latest_m15['uoedt']:.2f}]). "
                f"A Consolidated Trend is not active and EDT Stochastic is UNAVAILABLE."
            )
        elif discrete_state_code == "MCD3_BULL_VALUE":
            stoch_val = stoch if stoch is not None else 0.0
            return (
                f"XAUUSD has confirmed a BULLISH_CONSOLIDATED_TREND: M15 and M5 slopes are both UPTREND "
                f"({m15_sign}{m15_angle:.2f}° / {m5_sign}{m5_angle:.2f}°), M5 corridor nesting within M15 corridor is "
                f"{nesting_pct:.1f}% (>= 75.0% threshold), and current M5 corridor is totally engulfed inside M15 corridor. "
                f"Gold price is strongly confirmed to be under persistent bullish channel governance. "
                f"Standard EDT Stochastic is {stoch_val:.2f}% (Oversold / Value Dip Zone near LOEDT)."
            )
        elif discrete_state_code == "MCD3_BULL_MID":
            stoch_val = stoch if stoch is not None else 50.0
            return (
                f"XAUUSD has confirmed a BULLISH_CONSOLIDATED_TREND: M15 and M5 slopes are both UPTREND "
                f"({m15_sign}{m15_angle:.2f}° / {m5_sign}{m5_angle:.2f}°), M5 corridor nesting within M15 corridor is "
                f"{nesting_pct:.1f}% (>= 75.0% threshold), and current M5 corridor is totally engulfed inside M15 corridor. "
                f"Gold price is strongly confirmed to be under persistent bullish channel governance. "
                f"Standard EDT Stochastic is {stoch_val:.2f}% (Equilibrium Sweet Spot)."
            )
        elif discrete_state_code == "MCD3_BULL_TOP":
            stoch_val = stoch if stoch is not None else 100.0
            return (
                f"XAUUSD has confirmed a BULLISH_CONSOLIDATED_TREND: M15 and M5 slopes are both UPTREND "
                f"({m15_sign}{m15_angle:.2f}° / {m5_sign}{m5_angle:.2f}°), M5 corridor nesting within M15 corridor is "
                f"{nesting_pct:.1f}% (>= 75.0% threshold), and current M5 corridor is totally engulfed inside M15 corridor. "
                f"Gold price is strongly confirmed to be under persistent bullish channel governance. "
                f"Standard EDT Stochastic is {stoch_val:.2f}% (Overbought / Climax Zone near UOEDT)."
            )
        elif discrete_state_code == "MCD3_BEAR_PREMIUM":
            stoch_val = stoch if stoch is not None else 100.0
            return (
                f"XAUUSD has confirmed a BEARISH_CONSOLIDATED_TREND: M15 and M5 slopes are both DOWNTREND "
                f"({m15_sign}{m15_angle:.2f}° / {m5_sign}{m5_angle:.2f}°), M5 corridor nesting within M15 corridor is "
                f"{nesting_pct:.1f}% (>= 75.0% threshold), and current M5 corridor is totally engulfed inside M15 corridor. "
                f"Gold price is strongly confirmed to be under persistent bearish channel governance. "
                f"Standard EDT Stochastic is {stoch_val:.2f}% (Premium / Short Opportunity Zone near UOEDT)."
            )
        elif discrete_state_code == "MCD3_BEAR_MID":
            stoch_val = stoch if stoch is not None else 50.0
            return (
                f"XAUUSD has confirmed a BEARISH_CONSOLIDATED_TREND: M15 and M5 slopes are both DOWNTREND "
                f"({m15_sign}{m15_angle:.2f}° / {m5_sign}{m5_angle:.2f}°), M5 corridor nesting within M15 corridor is "
                f"{nesting_pct:.1f}% (>= 75.0% threshold), and current M5 corridor is totally engulfed inside M15 corridor. "
                f"Gold price is strongly confirmed to be under persistent bearish channel governance. "
                f"Standard EDT Stochastic is {stoch_val:.2f}% (Equilibrium Sweet Spot)."
            )
        elif discrete_state_code == "MCD3_BEAR_BOTTOM":
            stoch_val = stoch if stoch is not None else 0.0
            return (
                f"XAUUSD has confirmed a BEARISH_CONSOLIDATED_TREND: M15 and M5 slopes are both DOWNTREND "
                f"({m15_sign}{m15_angle:.2f}° / {m5_sign}{m5_angle:.2f}°), M5 corridor nesting within M15 corridor is "
                f"{nesting_pct:.1f}% (>= 75.0% threshold), and current M5 corridor is totally engulfed inside M15 corridor. "
                f"Gold price is strongly confirmed to be under persistent bearish channel governance. "
                f"Standard EDT Stochastic is {stoch_val:.2f}% (Oversold / Floor Caution Zone near LOEDT)."
            )
        elif discrete_state_code == "MCD3_SIDEWAYS_EQUILIBRIUM":
            stoch_val = stoch if stoch is not None else 50.0
            return (
                f"XAUUSD has confirmed a SIDEWAYS_CONSOLIDATED_TREND: M15 and M5 slopes are both horizontal SIDEWAYS "
                f"({m15_sign}{m15_angle:.2f}° / {m5_sign}{m5_angle:.2f}°), M5 corridor nesting within M15 corridor is "
                f"{nesting_pct:.1f}%, and current M5 corridor is engulfed inside M15 corridor. "
                f"Standard EDT Stochastic is {stoch_val:.2f}% (Range Equilibrium)."
            )
        else:
            return (
                f"XAUUSD multi-timeframe state is {discrete_state_code}. Evaluator executed cleanly."
            )

    def _build_invalid_payload(self, error_message: str) -> Dict[str, Any]:
        """Build safe INVALID payload complying with DavinTrade Stack D standard."""
        return {
            "symbol": self.symbol,
            "timeframe": "M15_M5",
            "timestamp": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
            "evaluator_version": "1.0.0",
            "module": "MCD3_CONSOLIDATED_TREND_AND_EDT_STOCHASTIC",
            "validation": {
                "status": "FAIL",
                "tier1_active_indicators": {
                    "m15": self.validation_checks.get("m15_active_detected", []),
                    "m5": self.validation_checks.get("m5_active_detected", []),
                },
                "errors": self.validation_errors + [error_message],
                "warnings": self.validation_warnings,
            },
            "m15_metrics": None,
            "m5_metrics": None,
            "consolidated_trend_conditions": {
                "condition_1_trend_aligned": False,
                "condition_2_historical_nesting_passed": False,
                "condition_2_nesting_rate_pct": 0.0,
                "condition_2_contained_bars": 0,
                "condition_2_total_bars": 0,
                "condition_3_current_bar_engulfed": False,
            },
            "evaluation": {
                "is_consolidated_trend": False,
                "consolidated_trend_state": "INVALID",
                "edt_stochastic": None,
                "stochastic_zone": "UNAVAILABLE",
                "regime_status": "UNCERTAIN",
                "discrete_state_code": "MCD3_INVALID",
                "tactical_bias": "NEUTRAL_STAND_ASIDE",
            },
            "commentary": f"MCD3 evaluation aborted due to pre-flight validation failure: {error_message}",
        }

    def run_cli(self, output_path: Optional[str] = None) -> Dict[str, Any]:
        """Run CLI evaluation, display formatted summary, and optionally save JSON."""
        print("=" * 80)
        print(" DavinTrade Stack D | Engine 1.5A Module")
        print(" MCD3: Consolidated Trend and EDT Stochastic Evaluator (XAUUSD - M15 & M5)")
        print("=" * 80)
        print(f" Excel Workbook: {self.excel_path}")
        print(f" Target M15 Indicator Override: {self.m15_target_indicator or 'None (Strict Production Mode)'}")
        print(f" Target M5 Indicator Override:  {self.m5_target_indicator or 'None (Strict Production Mode)'}")
        print("-" * 80)

        result = self.evaluate()

        val = result["validation"]
        eval_res = result["evaluation"]
        conds = result["consolidated_trend_conditions"]

        print(f" Validation Status:        {val['status']}")
        if val["status"] == "FAIL":
            print(f" Validation Errors:        {val.get('errors')}")
            print(f" Discrete State Code:      {eval_res['discrete_state_code']}")
            print(f" Commentary:               {result['commentary']}")
        else:
            m15 = result["m15_metrics"]
            m5 = result["m5_metrics"]
            print(f" M15 Active Indicator:     {m15['active_indicator']} ({m15['trend_state']} @ {m15['regression_angle']}°)")
            print(f" M5 Active Indicator:      {m5['active_indicator']} ({m5['trend_state']} @ {m5['regression_angle']}°)")
            print(f" Condition 1 (Aligned):    {conds['condition_1_trend_aligned']}")
            print(f" Condition 2 (Nesting):    {conds['condition_2_historical_nesting_passed']} ({conds['condition_2_nesting_rate_pct']}% - {conds['condition_2_contained_bars']}/{conds['condition_2_total_bars']} bars)")
            print(f" Condition 3 (Engulfment): {conds['condition_3_current_bar_engulfed']}")
            print(f" Consolidated Trend:       {eval_res['is_consolidated_trend']} ({eval_res['consolidated_trend_state']})")
            print(f" EDT Stochastic:           {eval_res['edt_stochastic'] if eval_res['edt_stochastic'] is not None else 'UNAVAILABLE (null)'}")
            print(f" Stochastic Zone:          {eval_res['stochastic_zone']}")
            print(f" Regime Status:            {eval_res['regime_status']}")
            print(f" Discrete State Code:      {eval_res['discrete_state_code']}")
            print(f" Tactical Bias:            {eval_res['tactical_bias']}")
            print("-" * 80)
            print(f" Commentary: {result['commentary']}")

        print("=" * 80)

        if output_path:
            out_file = os.path.abspath(output_path)
            os.makedirs(os.path.dirname(out_file), exist_ok=True)
            with open(out_file, "w", encoding="utf-8") as f:
                json.dump(result, f, indent=2, ensure_ascii=False)
            print(f" JSONB payload saved successfully to: {out_file}\n")

        return result


if __name__ == "__main__":
    base_dir = os.path.dirname(os.path.abspath(__file__))
    default_excel = os.path.abspath(os.path.join(base_dir, "..", "market_data_v6_replicated.xlsx"))
    default_output = os.path.abspath(os.path.join(base_dir, "mcd3_output.json"))

    # Support CLI arguments: [excel_path] [output_path] [--m15-ind NAME] [--m5-ind NAME]
    excel_arg = default_excel
    output_arg = default_output
    m15_ind_arg = None
    m5_ind_arg = None

    args = sys.argv[1:]
    i = 0
    pos_args = []
    while i < len(args):
        if args[i] == "--m15-ind" and i + 1 < len(args):
            m15_ind_arg = args[i + 1]
            i += 2
        elif args[i] == "--m5-ind" and i + 1 < len(args):
            m5_ind_arg = args[i + 1]
            i += 2
        elif args[i] in ("--output", "-o") and i + 1 < len(args):
            output_arg = args[i + 1]
            i += 2
        elif args[i] in ("--excel", "-e") and i + 1 < len(args):
            excel_arg = args[i + 1]
            i += 2
        else:
            pos_args.append(args[i])
            i += 1

    if len(pos_args) >= 1:
        excel_arg = pos_args[0]
    if len(pos_args) >= 2:
        output_arg = pos_args[1]

    evaluator = MCD3ConsolidatedTrendEvaluator(
        excel_path=excel_arg,
        m15_target_indicator=m15_ind_arg,
        m5_target_indicator=m5_ind_arg,
    )
    evaluator.run_cli(output_path=output_arg)
