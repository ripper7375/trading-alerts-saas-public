"""
Unit Tests for MCD3: Consolidated Trend and EDT Stochastic Evaluator
Tests Dual 4-Tier Pre-Flight Validation, the 3 Strict Conditions of Consolidated Trend,
Conditional Standard EDT Stochastic execution, and the full 10-state discrete synthesis matrix.
"""

import os
import sys
import tempfile
import unittest
import openpyxl

# Ensure UTF-8 output on Windows console
if sys.stdout and hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

base_dir = os.path.dirname(os.path.abspath(__file__))
if base_dir not in sys.path:
    sys.path.insert(0, base_dir)

from mcd3_evaluator import (
    MCD3ConsolidatedTrendEvaluator,
    MCD3ValidationError,
    M15_CANDIDATE_INDICATORS,
    M5_CANDIDATE_INDICATORS,
)


def create_synthetic_workbook(
    file_path: str,
    m15_active: str = "best_fit_a",
    m5_active: str = "best_fit_a",
    m15_angle: float = 10.0,
    m5_angle: float = 8.0,
    m15_cr: float = 70.0,
    m5_cr: float = 65.0,
    m15_corridor: tuple = (4300.0, 4400.0),  # (loedt, uoedt)
    m5_corridor: tuple = (4320.0, 4380.0),   # (loedt, uoedt)
    latest_m15_ssa: float = 4350.0,
    latest_m5_uoedt: float = 4380.0,
    latest_m5_loedt: float = 4320.0,
    nesting_override: float = 1.0,           # Fraction of historical bars nested
    m15_channel_inverted: bool = False,
    num_m15_bars: int = 100,
    num_m5_bars: int = 300,
):
    """Helper to generate a clean synthetic workbook for rigorous testing."""
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # 1. market_data_v6_M15
    ws_m15 = wb.create_sheet(title="market_data_v6_M15")
    m15_headers = ["id", "terminal_id", "timestamp", "symbol", "timeframe", "open", "high", "low", "close", "volume"]
    m15_headers.extend([f"{m15_active}_ssa", f"{m15_active}_uoedt", f"{m15_active}_loedt", f"{m15_active}_base_fl"])
    ws_m15.append(m15_headers)

    start_ts = 1780000000
    for i in range(num_m15_bars):
        ts = start_ts + i * 900
        close = 4350.0
        ssa = latest_m15_ssa if i == num_m15_bars - 1 else 4350.0
        uoedt = m15_corridor[0] if m15_channel_inverted else m15_corridor[1]
        loedt = m15_corridor[1] if m15_channel_inverted else m15_corridor[0]
        base_fl = 4350.0
        row = [f"id_m15_{i}", "worker", ts, "XAUUSD", "M15", close, close, close, close, 100, ssa, uoedt, loedt, base_fl]
        ws_m15.append(row)

    # 2. market_data_v6_M5
    ws_m5 = wb.create_sheet(title="market_data_v6_M5")
    m5_headers = ["id", "terminal_id", "timestamp", "symbol", "timeframe", "open", "high", "low", "close", "volume"]
    if m5_active == "fractal":
        m5_headers.extend(["fractal_best_fl", "fractal_uoedt", "fractal_loedt"])
    else:
        m5_headers.extend([f"{m5_active}_ssa", f"{m5_active}_uoedt", f"{m5_active}_loedt", f"{m5_active}_base_fl"])
    ws_m5.append(m5_headers)

    for i in range(num_m5_bars):
        ts = start_ts + i * 300
        close = 4350.0
        val = 4350.0

        if i == num_m5_bars - 1:
            uoedt = latest_m5_uoedt
            loedt = latest_m5_loedt
        else:
            # Check if this bar should be nested based on nesting_override fraction
            if i < int(num_m5_bars * nesting_override):
                uoedt = m5_corridor[1]
                loedt = m5_corridor[0]
            else:
                # Overflow outside M15 corridor
                uoedt = m15_corridor[1] + 50.0
                loedt = m15_corridor[0] - 50.0

        base_fl = 4350.0
        if m5_active == "fractal":
            row = [f"id_m5_{i}", "worker", ts, "XAUUSD", "M5", close, close, close, close, 100, base_fl, uoedt, loedt]
        else:
            row = [f"id_m5_{i}", "worker", ts, "XAUUSD", "M5", close, close, close, close, 100, val, uoedt, loedt, base_fl]
        ws_m5.append(row)

    # 3. indicator_statistics
    ws_stat = wb.create_sheet(title="indicator_statistics")
    stat_headers = ["symbol", "timeframe", "source", "regression_angle", "containment_rate", "containment_n"]
    ws_stat.append(stat_headers)

    ws_stat.append(["XAUUSD", "M15", m15_active, m15_angle, m15_cr, num_m15_bars])
    m5_stat_source = "fractal_edt" if m5_active == "fractal" else m5_active
    ws_stat.append(["XAUUSD", "M5", m5_stat_source, m5_angle, m5_cr, num_m5_bars])

    wb.save(file_path)


class TestMCD3ConsolidatedTrendLogic(unittest.TestCase):
    """Comprehensive test suite for MCD3 Evaluator."""

    @classmethod
    def setUpClass(cls):
        base_dir = os.path.dirname(os.path.abspath(__file__))
        cls.real_excel_path = os.path.abspath(os.path.join(base_dir, "..", "market_data_v6_replicated.xlsx"))
        cls.temp_dir = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)

    @classmethod
    def tearDownClass(cls):
        try:
            cls.temp_dir.cleanup()
        except Exception:
            pass

    def test_01_real_data_execution_non_consolidated(self):
        """Test real data execution with non_b (M15) and best_fit_a (M5) on market_data_v6_replicated.xlsx."""
        evaluator = MCD3ConsolidatedTrendEvaluator(
            excel_path=self.real_excel_path,
            m5_target_indicator="best_fit_a",
        )
        res = evaluator.evaluate()

        self.assertEqual(res["module"], "MCD3_CONSOLIDATED_TREND_AND_EDT_STOCHASTIC")
        self.assertEqual(res["validation"]["status"], "PASS")

        # M15 checks
        m15 = res["m15_metrics"]
        self.assertEqual(m15["active_indicator"], "non_b")
        self.assertEqual(m15["trend_state"], "DOWNTREND")
        self.assertEqual(m15["regression_angle"], -29.72)

        # M5 checks
        m5 = res["m5_metrics"]
        self.assertEqual(m5["active_indicator"], "best_fit_a")
        self.assertEqual(m5["trend_state"], "UPTREND")
        self.assertEqual(m5["regression_angle"], 6.94)

        # Conditions
        conds = res["consolidated_trend_conditions"]
        self.assertFalse(conds["condition_1_trend_aligned"])
        self.assertFalse(conds["condition_2_historical_nesting_passed"])
        self.assertFalse(conds["condition_3_current_bar_engulfed"])

        # Evaluation outcome
        eval_res = res["evaluation"]
        self.assertFalse(eval_res["is_consolidated_trend"])
        self.assertEqual(eval_res["consolidated_trend_state"], "NON_CONSOLIDATED")
        self.assertIsNone(eval_res["edt_stochastic"])
        self.assertEqual(eval_res["stochastic_zone"], "UNAVAILABLE")
        self.assertEqual(eval_res["regime_status"], "TREND_MISALIGNMENT")
        self.assertEqual(eval_res["discrete_state_code"], "MCD3_NON_CONSOLIDATED_TREND_CONFLICT")
        self.assertEqual(eval_res["tactical_bias"], "NEUTRAL_STAND_ASIDE")
        self.assertIn("M15 slope is DOWNTREND (-29.72°) while M5 slope is UPTREND (+6.94°)", res["commentary"])

    def test_02_synthetic_bullish_consolidated_value_zone(self):
        """Test Bullish Consolidated with Stochastic <= 20% -> MCD3_BULL_VALUE."""
        file_path = os.path.join(self.temp_dir.name, "bull_value.xlsx")
        create_synthetic_workbook(
            file_path=file_path,
            m15_angle=12.0,  # UPTREND
            m5_angle=8.0,    # UPTREND
            m15_corridor=(4300.0, 4400.0),
            m5_corridor=(4320.0, 4380.0),
            latest_m15_ssa=4310.0,  # (4310 - 4300) / 100 * 100 = 10.0%
        )

        evaluator = MCD3ConsolidatedTrendEvaluator(excel_path=file_path)
        res = evaluator.evaluate()

        self.assertEqual(res["validation"]["status"], "PASS")
        self.assertTrue(res["evaluation"]["is_consolidated_trend"])
        self.assertEqual(res["evaluation"]["consolidated_trend_state"], "BULLISH_CONSOLIDATED")
        self.assertEqual(res["evaluation"]["edt_stochastic"], 10.0)
        self.assertEqual(res["evaluation"]["stochastic_zone"], "VALUE_ZONE")
        self.assertEqual(res["evaluation"]["regime_status"], "BULLISH_CONSOLIDATED_VALUE_ZONE")
        self.assertEqual(res["evaluation"]["discrete_state_code"], "MCD3_BULL_VALUE")
        self.assertEqual(res["evaluation"]["tactical_bias"], "HIGH_CONVICTION_BUY_DIP")
        self.assertIn("Oversold / Value Dip Zone near LOEDT", res["commentary"])

    def test_03_synthetic_bullish_consolidated_equilibrium(self):
        """Test Bullish Consolidated with Stochastic 50% -> MCD3_BULL_MID."""
        file_path = os.path.join(self.temp_dir.name, "bull_mid.xlsx")
        create_synthetic_workbook(
            file_path=file_path,
            m15_angle=12.0,
            m5_angle=8.0,
            latest_m15_ssa=4350.0,  # 50.0%
        )

        evaluator = MCD3ConsolidatedTrendEvaluator(excel_path=file_path)
        res = evaluator.evaluate()

        self.assertEqual(res["evaluation"]["edt_stochastic"], 50.0)
        self.assertEqual(res["evaluation"]["stochastic_zone"], "EQUILIBRIUM")
        self.assertEqual(res["evaluation"]["regime_status"], "BULLISH_CONSOLIDATED_EQUILIBRIUM")
        self.assertEqual(res["evaluation"]["discrete_state_code"], "MCD3_BULL_MID")
        self.assertEqual(res["evaluation"]["tactical_bias"], "HOLD_BULLISH_TREND_RUNNER")
        self.assertIn("Equilibrium Sweet Spot", res["commentary"])

    def test_04_synthetic_bullish_consolidated_overbought(self):
        """Test Bullish Consolidated with Stochastic >= 80% -> MCD3_BULL_TOP."""
        file_path = os.path.join(self.temp_dir.name, "bull_top.xlsx")
        create_synthetic_workbook(
            file_path=file_path,
            m15_angle=12.0,
            m5_angle=8.0,
            latest_m15_ssa=4390.0,  # 90.0%
        )

        evaluator = MCD3ConsolidatedTrendEvaluator(excel_path=file_path)
        res = evaluator.evaluate()

        self.assertEqual(res["evaluation"]["edt_stochastic"], 90.0)
        self.assertEqual(res["evaluation"]["stochastic_zone"], "OVERBOUGHT_ZONE")
        self.assertEqual(res["evaluation"]["regime_status"], "BULLISH_CONSOLIDATED_OVERBOUGHT")
        self.assertEqual(res["evaluation"]["discrete_state_code"], "MCD3_BULL_TOP")
        self.assertEqual(res["evaluation"]["tactical_bias"], "CAUTION_TAKE_PROFIT_BUY")
        self.assertIn("Overbought / Climax Zone near UOEDT", res["commentary"])

    def test_05_synthetic_bearish_consolidated_premium(self):
        """Test Bearish Consolidated with Stochastic >= 80% -> MCD3_BEAR_PREMIUM."""
        file_path = os.path.join(self.temp_dir.name, "bear_premium.xlsx")
        create_synthetic_workbook(
            file_path=file_path,
            m15_angle=-15.0,  # DOWNTREND
            m5_angle=-10.0,   # DOWNTREND
            latest_m15_ssa=4390.0,  # 90.0%
        )

        evaluator = MCD3ConsolidatedTrendEvaluator(excel_path=file_path)
        res = evaluator.evaluate()

        self.assertTrue(res["evaluation"]["is_consolidated_trend"])
        self.assertEqual(res["evaluation"]["consolidated_trend_state"], "BEARISH_CONSOLIDATED")
        self.assertEqual(res["evaluation"]["edt_stochastic"], 90.0)
        self.assertEqual(res["evaluation"]["stochastic_zone"], "PREMIUM_ZONE")
        self.assertEqual(res["evaluation"]["regime_status"], "BEARISH_CONSOLIDATED_PREMIUM_ZONE")
        self.assertEqual(res["evaluation"]["discrete_state_code"], "MCD3_BEAR_PREMIUM")
        self.assertEqual(res["evaluation"]["tactical_bias"], "HIGH_CONVICTION_SELL_RALLY")
        self.assertIn("Premium / Short Opportunity Zone near UOEDT", res["commentary"])

    def test_06_synthetic_bearish_consolidated_equilibrium(self):
        """Test Bearish Consolidated with Stochastic 50% -> MCD3_BEAR_MID."""
        file_path = os.path.join(self.temp_dir.name, "bear_mid.xlsx")
        create_synthetic_workbook(
            file_path=file_path,
            m15_angle=-15.0,
            m5_angle=-10.0,
            latest_m15_ssa=4350.0,  # 50.0%
        )

        evaluator = MCD3ConsolidatedTrendEvaluator(excel_path=file_path)
        res = evaluator.evaluate()

        self.assertEqual(res["evaluation"]["edt_stochastic"], 50.0)
        self.assertEqual(res["evaluation"]["stochastic_zone"], "EQUILIBRIUM")
        self.assertEqual(res["evaluation"]["regime_status"], "BEARISH_CONSOLIDATED_EQUILIBRIUM")
        self.assertEqual(res["evaluation"]["discrete_state_code"], "MCD3_BEAR_MID")
        self.assertEqual(res["evaluation"]["tactical_bias"], "HOLD_BEARISH_TREND_RUNNER")

    def test_07_synthetic_bearish_consolidated_oversold(self):
        """Test Bearish Consolidated with Stochastic <= 20% -> MCD3_BEAR_BOTTOM."""
        file_path = os.path.join(self.temp_dir.name, "bear_bottom.xlsx")
        create_synthetic_workbook(
            file_path=file_path,
            m15_angle=-15.0,
            m5_angle=-10.0,
            latest_m15_ssa=4310.0,  # 10.0%
        )

        evaluator = MCD3ConsolidatedTrendEvaluator(excel_path=file_path)
        res = evaluator.evaluate()

        self.assertEqual(res["evaluation"]["edt_stochastic"], 10.0)
        self.assertEqual(res["evaluation"]["stochastic_zone"], "OVERSOLD_ZONE")
        self.assertEqual(res["evaluation"]["regime_status"], "BEARISH_CONSOLIDATED_OVERSOLD")
        self.assertEqual(res["evaluation"]["discrete_state_code"], "MCD3_BEAR_BOTTOM")
        self.assertEqual(res["evaluation"]["tactical_bias"], "CAUTION_TAKE_PROFIT_SELL")
        self.assertIn("Oversold / Floor Caution Zone near LOEDT", res["commentary"])

    def test_08_synthetic_sideways_consolidated(self):
        """Test Sideways Consolidated -> MCD3_SIDEWAYS_EQUILIBRIUM."""
        file_path = os.path.join(self.temp_dir.name, "sideways.xlsx")
        create_synthetic_workbook(
            file_path=file_path,
            m15_angle=2.0,   # SIDEWAYS
            m5_angle=-1.5,   # SIDEWAYS
            latest_m15_ssa=4350.0,
        )

        evaluator = MCD3ConsolidatedTrendEvaluator(excel_path=file_path)
        res = evaluator.evaluate()

        self.assertTrue(res["evaluation"]["is_consolidated_trend"])
        self.assertEqual(res["evaluation"]["consolidated_trend_state"], "SIDEWAYS_CONSOLIDATED")
        self.assertEqual(res["evaluation"]["stochastic_zone"], "IN_CORRIDOR")
        self.assertEqual(res["evaluation"]["regime_status"], "SIDEWAYS_CONSOLIDATED_EQUILIBRIUM")
        self.assertEqual(res["evaluation"]["discrete_state_code"], "MCD3_SIDEWAYS_EQUILIBRIUM")
        self.assertEqual(res["evaluation"]["tactical_bias"], "RANGE_BOUND_MEAN_REVERSION")

    def test_09_condition1_trend_conflict_failure(self):
        """Test Condition 1 Failure: M15 Uptrend vs M5 Downtrend -> MCD3_NON_CONSOLIDATED_TREND_CONFLICT."""
        file_path = os.path.join(self.temp_dir.name, "cond1_fail.xlsx")
        create_synthetic_workbook(
            file_path=file_path,
            m15_angle=15.0,  # UPTREND
            m5_angle=-12.0,  # DOWNTREND
        )

        evaluator = MCD3ConsolidatedTrendEvaluator(excel_path=file_path)
        res = evaluator.evaluate()

        self.assertFalse(res["consolidated_trend_conditions"]["condition_1_trend_aligned"])
        self.assertFalse(res["evaluation"]["is_consolidated_trend"])
        self.assertIsNone(res["evaluation"]["edt_stochastic"])
        self.assertEqual(res["evaluation"]["regime_status"], "TREND_MISALIGNMENT")
        self.assertEqual(res["evaluation"]["discrete_state_code"], "MCD3_NON_CONSOLIDATED_TREND_CONFLICT")

    def test_10_condition2_insufficient_nesting_failure(self):
        """Test Condition 2 Failure: Nesting rate < 75% -> MCD3_NON_CONSOLIDATED_OVERFLOW."""
        file_path = os.path.join(self.temp_dir.name, "cond2_fail.xlsx")
        create_synthetic_workbook(
            file_path=file_path,
            m15_angle=10.0,
            m5_angle=8.0,
            nesting_override=0.50,  # Only 50% nested (< 75%)
        )

        evaluator = MCD3ConsolidatedTrendEvaluator(excel_path=file_path)
        res = evaluator.evaluate()

        self.assertTrue(res["consolidated_trend_conditions"]["condition_1_trend_aligned"])
        self.assertFalse(res["consolidated_trend_conditions"]["condition_2_historical_nesting_passed"])
        self.assertFalse(res["evaluation"]["is_consolidated_trend"])
        self.assertIsNone(res["evaluation"]["edt_stochastic"])
        self.assertEqual(res["evaluation"]["regime_status"], "INSUFFICIENT_CORRIDOR_NESTING")
        self.assertEqual(res["evaluation"]["discrete_state_code"], "MCD3_NON_CONSOLIDATED_OVERFLOW")
        self.assertIn("fails the strict 75.0% threshold requirement", res["commentary"])

    def test_11_condition3_current_bar_escape_failure(self):
        """Test Condition 3 Failure: Current bar M5 corridor breached outside M15 -> MCD3_NON_CONSOLIDATED_ESCAPE."""
        file_path = os.path.join(self.temp_dir.name, "cond3_fail.xlsx")
        create_synthetic_workbook(
            file_path=file_path,
            m15_angle=10.0,
            m5_angle=8.0,
            m15_corridor=(4300.0, 4400.0),
            m5_corridor=(4320.0, 4380.0),
            nesting_override=1.0,
            latest_m5_uoedt=4410.0,  # Breaches outside M15 UOEDT (4400.0)
            latest_m5_loedt=4320.0,
        )

        evaluator = MCD3ConsolidatedTrendEvaluator(excel_path=file_path)
        res = evaluator.evaluate()

        self.assertTrue(res["consolidated_trend_conditions"]["condition_1_trend_aligned"])
        self.assertTrue(res["consolidated_trend_conditions"]["condition_2_historical_nesting_passed"])
        self.assertFalse(res["consolidated_trend_conditions"]["condition_3_current_bar_engulfed"])
        self.assertFalse(res["evaluation"]["is_consolidated_trend"])
        self.assertIsNone(res["evaluation"]["edt_stochastic"])
        self.assertEqual(res["evaluation"]["regime_status"], "CURRENT_CORRIDOR_ESCAPE")
        self.assertEqual(res["evaluation"]["discrete_state_code"], "MCD3_NON_CONSOLIDATED_ESCAPE")
        self.assertIn("current bar extends outside the M15 corridor boundaries", res["commentary"])

    def test_12_tier1_strict_production_multi_and_zero_indicator_errors(self):
        """Test Tier 1: Multiple active indicators and zero active indicators trigger INVALID."""
        # 1. Real workbook without target override has multiple indicators on M5
        evaluator_multi = MCD3ConsolidatedTrendEvaluator(excel_path=self.real_excel_path)
        res_multi = evaluator_multi.evaluate()
        self.assertEqual(res_multi["validation"]["status"], "FAIL")
        self.assertEqual(res_multi["evaluation"]["discrete_state_code"], "MCD3_INVALID")
        self.assertIn("Multiple active indicators detected on M5", res_multi["validation"]["errors"][0])

        # 2. Synthetic workbook with empty columns (zero active indicators)
        file_path = os.path.join(self.temp_dir.name, "zero_ind.xlsx")
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        ws_m15 = wb.create_sheet(title="market_data_v6_M15")
        ws_m15.append(["timestamp", "close"])
        for i in range(100):
            ws_m15.append([1780000000 + i * 900, 4350.0])
        ws_m5 = wb.create_sheet(title="market_data_v6_M5")
        ws_m5.append(["timestamp", "close"])
        for i in range(100):
            ws_m5.append([1780000000 + i * 300, 4350.0])
        ws_stat = wb.create_sheet(title="indicator_statistics")
        ws_stat.append(["symbol", "timeframe", "source", "regression_angle", "containment_rate"])
        wb.save(file_path)

        evaluator_zero = MCD3ConsolidatedTrendEvaluator(excel_path=file_path)
        res_zero = evaluator_zero.evaluate()
        self.assertEqual(res_zero["validation"]["status"], "FAIL")
        self.assertEqual(res_zero["evaluation"]["discrete_state_code"], "MCD3_INVALID")

    def test_13_tier3_and_tier4_channel_and_stats_failures(self):
        """Test Tier 3 (Inverted channel) and Tier 4 (Low containment rate < 50%) failures."""
        # 1. Inverted channel on M15
        file_inverted = os.path.join(self.temp_dir.name, "inverted.xlsx")
        create_synthetic_workbook(file_path=file_inverted, m15_channel_inverted=True)
        evaluator_inv = MCD3ConsolidatedTrendEvaluator(excel_path=file_inverted)
        res_inv = evaluator_inv.evaluate()
        self.assertEqual(res_inv["validation"]["status"], "FAIL")
        self.assertEqual(res_inv["evaluation"]["discrete_state_code"], "MCD3_INVALID")
        self.assertTrue(any("Channel inverted" in err for err in res_inv["validation"]["errors"]))

        # 2. Low containment rate (< 50%)
        file_low_cr = os.path.join(self.temp_dir.name, "low_cr.xlsx")
        create_synthetic_workbook(file_path=file_low_cr, m15_cr=45.0)
        evaluator_cr = MCD3ConsolidatedTrendEvaluator(excel_path=file_low_cr)
        res_cr = evaluator_cr.evaluate()
        self.assertEqual(res_cr["validation"]["status"], "FAIL")
        self.assertEqual(res_cr["evaluation"]["discrete_state_code"], "MCD3_INVALID")
        self.assertTrue(any("below the minimum threshold" in err for err in res_cr["validation"]["errors"]))


if __name__ == "__main__":
    unittest.main()
