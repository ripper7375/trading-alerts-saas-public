"""
Script to create an Excel file replicating the Prisma schema tables 'market_data_v6'
and 'indicator_statistics', populating it with all data from engine-1-5-new.
100% compliant with 2026-09-20 Prisma Schema (MarketDataV6: 98 fields, IndicatorStatistic: 88 fields).
"""

import os
import sys
import json
import hashlib
from datetime import datetime, timezone
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = r"d:\SaaS Project\trading-alerts-saas-public\davintrade-stack-d-and-e"
DATA_DIR = os.path.join(BASE_DIR, "engine-1-5-new")
OUTPUT_FILE_1 = os.path.join(DATA_DIR, "market_data_v6_replicated.xlsx")
OUTPUT_FILE_2 = os.path.join(BASE_DIR, "market_data_v6_replicated.xlsx")

# 98 columns of MarketDataV6 in exact Prisma schema order
MARKET_DATA_V6_COLUMNS = [
    # 1. Identity & Base OHLCV Spine (10)
    "id",
    "terminal_id",
    "timestamp",
    "symbol",
    "timeframe",
    "open",
    "high",
    "low",
    "close",
    "volume",
    # 2. Centroid best_fit_a (8)
    "best_fit_a_horiz_high_map",
    "best_fit_a_horiz_low_map",
    "best_fit_a_ssa",
    "best_fit_a_ema_ssa",
    "best_fit_a_crossing",
    "best_fit_a_base_fl",
    "best_fit_a_uoedt",
    "best_fit_a_loedt",
    # 3. Centroid best_fit_b (8)
    "best_fit_b_horiz_high_map",
    "best_fit_b_horiz_low_map",
    "best_fit_b_ssa",
    "best_fit_b_ema_ssa",
    "best_fit_b_crossing",
    "best_fit_b_base_fl",
    "best_fit_b_uoedt",
    "best_fit_b_loedt",
    # 4. Centroid cherry_a (8)
    "cherry_a_horiz_high_map",
    "cherry_a_horiz_low_map",
    "cherry_a_ssa",
    "cherry_a_ema_ssa",
    "cherry_a_crossing",
    "cherry_a_base_fl",
    "cherry_a_uoedt",
    "cherry_a_loedt",
    # 5. Centroid cherry_b (8)
    "cherry_b_horiz_high_map",
    "cherry_b_horiz_low_map",
    "cherry_b_ssa",
    "cherry_b_ema_ssa",
    "cherry_b_crossing",
    "cherry_b_base_fl",
    "cherry_b_uoedt",
    "cherry_b_loedt",
    # 6. Centroid most_recent (8)
    "most_recent_horiz_high_map",
    "most_recent_horiz_low_map",
    "most_recent_ssa",
    "most_recent_ema_ssa",
    "most_recent_crossing",
    "most_recent_base_fl",
    "most_recent_uoedt",
    "most_recent_loedt",
    # 7. Centroid non_a (8)
    "non_a_horiz_high_map",
    "non_a_horiz_low_map",
    "non_a_ssa",
    "non_a_ema_ssa",
    "non_a_crossing",
    "non_a_base_fl",
    "non_a_uoedt",
    "non_a_loedt",
    # 8. Centroid non_b (8)
    "non_b_horiz_high_map",
    "non_b_horiz_low_map",
    "non_b_ssa",
    "non_b_ema_ssa",
    "non_b_crossing",
    "non_b_base_fl",
    "non_b_uoedt",
    "non_b_loedt",
    # 9. Fractal EDT + Single Best Lines (5)
    "fractal_best_fl",
    "fractal_uoedt",
    "fractal_loedt",
    "best_resistance",
    "best_support",
    # 10. Auto-calibrated Support & Resistance Levels sr_1..sr_8 (8)
    "sr_1",
    "sr_2",
    "sr_3",
    "sr_4",
    "sr_5",
    "sr_6",
    "sr_7",
    "sr_8",
    # 11. Z-Score Candle Volatility (3)
    "body_direction",
    "body_size",
    "body_classification",
    # 12. ZigZag Market Structure & Metrics (11)
    "zigzag_point_type",
    "zigzag_current_point",
    "zigzag_price_change",
    "zigzag_pct_change",
    "zigzag_pct_change_class",
    "zigzag_bars",
    "zigzag_bars_class",
    "zigzag_price_per_bar",
    "zigzag_price_per_bar_class",
    "zigzag_slope",
    "zigzag_category",
    # 13. Provenance (3)
    "cycle_id",
    "collected_at",
    "calculated_at",
    # 14. Timestamps (2)
    "createdAt",
    "updatedAt",
]

# 88 columns of IndicatorStatistic in exact Prisma schema order (2026-09-20)
INDICATOR_STATISTIC_COLUMNS = [
    # 1. Identity & Ingestion Timing (8)
    "id",
    "terminal_id",
    "symbol",
    "timeframe",
    "source",
    "captured_at",
    "live_bar_ts",
    "cycle_id",
    # 2. Resolved Line (6)
    "solution_found",
    "raw_slope",
    "anchored_y_int",
    "regression_angle",
    "line_origin_ts",
    "touches",
    # 3. Fitting Window (5)
    "window_start_ts",
    "window_end_ts",
    "window_bars",
    "math_lookback",
    "crossings_n",
    # 4. Model A Crossings (6)
    "model_a_n",
    "model_a_r2",
    "model_a_mse",
    "model_a_var_ratio",
    "model_a_skew",
    "model_a_kurt",
    # 5. Model B Close Price (6)
    "model_b_n",
    "model_b_r2",
    "model_b_mse",
    "model_b_var_ratio",
    "model_b_skew",
    "model_b_kurt",
    # 6. EDT Channel (5)
    "uoedt_offset",
    "loedt_offset",
    "containment_n",
    "containment_count",
    "containment_rate",
    # 7. Extended Fit Window (11)
    "window_span_bars",
    "visual_window_bars",
    "bars_available",
    "leftmost_bar_index",
    "line_span_bars",
    "baseline_coverage_n",
    "baseline_coverage_rate",
    "centroids_used",
    "crossings_in_window_n",
    "first_crossing_ts",
    "last_crossing_ts",
    # 8. Extended Price Context (11)
    "live_close",
    "baseline_value",
    "uoedt_value",
    "loedt_value",
    "dist_to_baseline",
    "dist_to_uoedt",
    "dist_to_loedt",
    "channel_position",
    "window_high",
    "window_low",
    "window_range",
    # 9. Extended Channel Geometry (6)
    "channel_width",
    "channel_asymmetry",
    "breach_above_n",
    "breach_below_n",
    "max_excursion_above",
    "max_excursion_below",
    # 10. Residual Diagnostics Crossings (6)
    "resid_a_n",
    "resid_a_mean",
    "resid_a_mae",
    "resid_a_sd",
    "resid_a_max",
    "resid_a_dw",
    # 11. Residual Diagnostics Close Price (6)
    "resid_b_n",
    "resid_b_mean",
    "resid_b_mae",
    "resid_b_sd",
    "resid_b_max",
    "resid_b_dw",
    # 12. Support-Resistance Calibration Provenance (10)
    "sr_fractals_n",
    "sr_q25",
    "sr_q75",
    "sr_iqr",
    "sr_optimal_step",
    "sr_macro_clusters",
    "sr_nearest_resistance",
    "sr_nearest_support",
    "sr_dist_resistance_pts",
    "sr_dist_support_pts",
    # 13. System Configuration & Timestamp (2)
    "config_hash",
    "createdAt",
]


def snap_ts(ts: int, tf: str) -> int:
    tf_sec = 300 if tf == "M5" else 900
    return round(ts / tf_sec) * tf_sec


def make_cuid(prefix: str, key_str: str) -> str:
    h = hashlib.sha256(key_str.encode("utf-8")).hexdigest()
    return f"{prefix}{h[:23]}"


def load_all_market_data():
    records = {}  # (tf, ts) -> dict

    # 1. Ingest OHLCV FIRST as the foundational spine (Exact 3,000 bars for M5, 3,000 bars for M15)
    for tf in ["M5", "M15"]:
        p = os.path.join(DATA_DIR, f"OHLCV_XAUUSD_{tf}.txt")
        if not os.path.exists(p):
            continue
        with open(p, "r", encoding="utf-8") as f:
            f.readline()
            for line in f:
                parts = line.strip().split("\t")
                if not parts or not parts[0]:
                    continue
                ts = snap_ts(int(float(parts[0])), tf)
                dt_iso = datetime.fromtimestamp(ts, tz=timezone.utc).strftime(
                    "%Y-%m-%d %H:%M:%S"
                )
                records[(tf, ts)] = {
                    "id": make_cuid("c", f"market_data_v6_XAUUSD_{tf}_{ts}"),
                    "terminal_id": "push_worker_v5",
                    "timestamp": ts,
                    "symbol": "XAUUSD",
                    "timeframe": tf,
                    "open": float(parts[4]),
                    "high": float(parts[5]),
                    "low": float(parts[6]),
                    "close": float(parts[3]),
                    "volume": int(float(parts[7])),
                    "cycle_id": 1,
                    "collected_at": 1789764900,
                    "calculated_at": 1789764900,
                    "createdAt": dt_iso,
                    "updatedAt": dt_iso,
                }

    def get_row(tf: str, ts: int):
        return records.get((tf, ts))

    # 2. Centriod_Best_Fit_A_XAUUSD_M5
    p = os.path.join(DATA_DIR, "Centriod_Best_Fit_A_XAUUSD_M5.txt")
    if os.path.exists(p):
        with open(p, "r", encoding="utf-8") as f:
            f.readline()
            for line in f:
                parts = line.strip().split("\t")
                if not parts or not parts[0]:
                    continue
                ts = snap_ts(int(float(parts[0])), "M5")
                r = get_row("M5", ts)
                if not r:
                    continue
                if len(parts) > 4 and parts[4].strip():
                    r["best_fit_a_base_fl"] = float(parts[4])
                if len(parts) > 5 and parts[5].strip():
                    r["best_fit_a_uoedt"] = float(parts[5])
                if len(parts) > 6 and parts[6].strip():
                    r["best_fit_a_loedt"] = float(parts[6])
                if len(parts) > 7 and parts[7].strip():
                    v = float(parts[7])
                    if v > 0:
                        r["best_fit_a_horiz_high_map"] = v
                if len(parts) > 8 and parts[8].strip():
                    v = float(parts[8])
                    if v > 0:
                        r["best_fit_a_horiz_low_map"] = v
                if len(parts) > 9 and parts[9].strip():
                    v = float(parts[9])
                    if v > 0:
                        r["best_fit_a_ssa"] = v
                if len(parts) > 10 and parts[10].strip():
                    v = float(parts[10])
                    if v > 0:
                        r["best_fit_a_ema_ssa"] = v
                if len(parts) > 11 and parts[11].strip():
                    r["best_fit_a_crossing"] = int(float(parts[11]))

    # 3. Non-Recent-B_XAUUSD_M15
    p = os.path.join(DATA_DIR, "Non-Recent-B_XAUUSD_M15.txt")
    if os.path.exists(p):
        with open(p, "r", encoding="utf-8") as f:
            f.readline()
            for line in f:
                parts = line.strip().split("\t")
                if not parts or not parts[0]:
                    continue
                ts = snap_ts(int(float(parts[0])), "M15")
                r = get_row("M15", ts)
                if not r:
                    continue
                if len(parts) > 4 and parts[4].strip():
                    r["non_b_base_fl"] = float(parts[4])
                if len(parts) > 5 and parts[5].strip():
                    r["non_b_uoedt"] = float(parts[5])
                if len(parts) > 6 and parts[6].strip():
                    r["non_b_loedt"] = float(parts[6])
                if len(parts) > 7 and parts[7].strip():
                    v = float(parts[7])
                    if v > 0:
                        r["non_b_horiz_high_map"] = v
                if len(parts) > 8 and parts[8].strip():
                    v = float(parts[8])
                    if v > 0:
                        r["non_b_horiz_low_map"] = v
                if len(parts) > 9 and parts[9].strip():
                    v = float(parts[9])
                    if v > 0:
                        r["non_b_ssa"] = v
                if len(parts) > 10 and parts[10].strip():
                    v = float(parts[10])
                    if v > 0:
                        r["non_b_ema_ssa"] = v
                if len(parts) > 11 and parts[11].strip():
                    r["non_b_crossing"] = int(float(parts[11]))

    # 4. Fractal_EDT_XAUUSD_M5
    p = os.path.join(DATA_DIR, "Fractal_EDT_XAUUSD_M5.txt")
    if os.path.exists(p):
        with open(p, "r", encoding="utf-8") as f:
            f.readline()
            for line in f:
                parts = line.strip().split("\t")
                if not parts or not parts[0]:
                    continue
                ts = snap_ts(int(float(parts[0])), "M5")
                r = get_row("M5", ts)
                if not r:
                    continue
                if len(parts) > 4 and parts[4].strip():
                    r["fractal_best_fl"] = float(parts[4])
                if len(parts) > 5 and parts[5].strip():
                    r["fractal_uoedt"] = float(parts[5])
                if len(parts) > 6 and parts[6].strip():
                    r["fractal_loedt"] = float(parts[6])

    # 5. Resistance_Line_XAUUSD_M5
    p = os.path.join(DATA_DIR, "Resistance_Line_XAUUSD_M5.txt")
    if os.path.exists(p):
        with open(p, "r", encoding="utf-8") as f:
            f.readline()
            for line in f:
                parts = line.strip().split("\t")
                if not parts or not parts[0]:
                    continue
                ts = snap_ts(int(float(parts[0])), "M5")
                r = get_row("M5", ts)
                if not r:
                    continue
                if len(parts) > 4 and parts[4].strip():
                    r["best_resistance"] = float(parts[4])

    # 6. Support_Line_XAUUSD_M5
    p = os.path.join(DATA_DIR, "Support_Line_XAUUSD_M5.txt")
    if os.path.exists(p):
        with open(p, "r", encoding="utf-8") as f:
            f.readline()
            for line in f:
                parts = line.strip().split("\t")
                if not parts or not parts[0]:
                    continue
                ts = snap_ts(int(float(parts[0])), "M5")
                r = get_row("M5", ts)
                if not r:
                    continue
                if len(parts) > 4 and parts[4].strip():
                    r["best_support"] = float(parts[4])

    # 7. SR_Levels_XAUUSD_M15 (14th Indicator: Freedman-Diaconis IQR Clustering)
    p = os.path.join(DATA_DIR, "SR_Levels_XAUUSD_M15.txt")
    if os.path.exists(p):
        with open(p, "r", encoding="utf-8") as f:
            f.readline()
            for line in f:
                parts = line.strip().split("\t")
                if not parts or not parts[0]:
                    continue
                ts = snap_ts(int(float(parts[0])), "M15")
                r = get_row("M15", ts)
                if not r:
                    continue
                for idx in range(1, 9):
                    col_i = 3 + idx
                    if len(parts) > col_i and parts[col_i].strip():
                        try:
                            val = float(parts[col_i])
                            if val > 0:
                                r[f"sr_{idx}"] = val
                        except ValueError:
                            pass

    # 8. ZScore
    for tf in ["M5", "M15"]:
        p = os.path.join(DATA_DIR, f"ZScore_XAUUSD_{tf}.txt")
        if not os.path.exists(p):
            continue
        with open(p, "r", encoding="utf-8") as f:
            f.readline()
            for line in f:
                parts = line.strip().split("\t")
                if not parts or not parts[0]:
                    continue
                ts = snap_ts(int(float(parts[0])), tf)
                r = get_row(tf, ts)
                if not r:
                    continue
                if len(parts) > 7 and parts[7].strip():
                    r["body_direction"] = int(float(parts[7]))
                if len(parts) > 8 and parts[8].strip():
                    r["body_size"] = float(parts[8])
                if len(parts) > 9 and parts[9].strip():
                    r["body_classification"] = int(float(parts[9]))

    # 9. ZigZag
    for tf in ["M5", "M15"]:
        p = os.path.join(DATA_DIR, f"ZigZag_XAUUSD_{tf}.txt")
        if not os.path.exists(p):
            continue
        with open(p, "r", encoding="utf-8") as f:
            f.readline()
            for line in f:
                parts = line.strip().split("\t")
                if not parts or not parts[0]:
                    continue
                ts = snap_ts(int(float(parts[0])), tf)
                r = get_row(tf, ts)
                if not r:
                    continue
                if len(parts) > 4 and parts[4].strip():
                    r["zigzag_point_type"] = parts[4].strip()
                if len(parts) > 5 and parts[5].strip():
                    r["zigzag_current_point"] = float(parts[5])
                if len(parts) > 6 and parts[6].strip():
                    r["zigzag_price_change"] = float(parts[6])
                if len(parts) > 7 and parts[7].strip():
                    r["zigzag_pct_change"] = float(parts[7])
                if len(parts) > 8 and parts[8].strip():
                    r["zigzag_pct_change_class"] = int(float(parts[8]))
                if len(parts) > 9 and parts[9].strip():
                    r["zigzag_bars"] = int(float(parts[9]))
                if len(parts) > 10 and parts[10].strip():
                    r["zigzag_bars_class"] = int(float(parts[10]))
                if len(parts) > 11 and parts[11].strip():
                    r["zigzag_price_per_bar"] = float(parts[11])
                if len(parts) > 12 and parts[12].strip():
                    r["zigzag_price_per_bar_class"] = int(float(parts[12]))
                if len(parts) > 13 and parts[13].strip():
                    r["zigzag_slope"] = float(parts[13])
                if len(parts) > 14 and parts[14].strip():
                    r["zigzag_category"] = parts[14].strip()

    return records


# Statistic parsing fields (Exact match to export_collector_validator_v2.py production pipeline)
STAT_FIELDS = [
    ("Raw Slope (b)", None, "raw_slope", "real"),
    ("Anchored Y-Int", None, "anchored_y_int", "real"),
    ("Regression Angle", None, "regression_angle", "real"),
    ("Solution Found", None, "solution_found", "bool"),
    ("Line Origin TS (UTC)", None, "line_origin_ts", "int"),
    ("Best FL Touches", None, "touches", "int"),
    ("Window Start TS (UTC)", None, "window_start_ts", "int"),
    ("Window End TS (UTC)", None, "window_end_ts", "int"),
    (
        (
            "Observation Bars",
            "Observation Window (Box B Bars)",
            "Observation Window (Bars)",
        ),
        None,
        "window_bars",
        "int",
    ),
    (
        ("Math Window Bars", "Math Search Window (Bars)", "Max Window Bars"),
        None,
        "math_lookback",
        "int",
    ),
    ("Total 171 Crossings (n)", None, "crossings_n", "int"),
    ("Sample (n)", "[MODEL A", "model_a_n", "int"),
    ("R-Square", "[MODEL A", "model_a_r2", "real"),
    ("MSE", "[MODEL A", "model_a_mse", "real"),
    ("Var Ratio", "[MODEL A", "model_a_var_ratio", "real"),
    ("Skewness", "[MODEL A", "model_a_skew", "real"),
    ("Kurtosis", "[MODEL A", "model_a_kurt", "real"),
    ("Sample (n)", "[MODEL B", "model_b_n", "int"),
    ("R-Square", "[MODEL B", "model_b_r2", "real"),
    ("MSE", "[MODEL B", "model_b_mse", "real"),
    ("Var Ratio", "[MODEL B", "model_b_var_ratio", "real"),
    ("Skewness", "[MODEL B", "model_b_skew", "real"),
    ("Kurtosis", "[MODEL B", "model_b_kurt", "real"),
    ("UOEDT Offset", "[EDT CHANNEL", "uoedt_offset", "real"),
    ("LOEDT Offset", "[EDT CHANNEL", "loedt_offset", "real"),
    ("Containment Sample (n)", "[EDT CHANNEL", "containment_n", "int"),
    ("Containment Count", "[EDT CHANNEL", "containment_count", "int"),
    ("Containment Rate", "[EDT CHANNEL", "containment_rate", "real"),
    # Extended statistics [added 2026-09-20]
    ("Window Bars", None, "window_span_bars", "int"),
    ("Visual Window Bars", "[FIT WINDOW", "visual_window_bars", "int"),
    ("Bars Available", "[FIT WINDOW", "bars_available", "int"),
    ("Leftmost Bar Index", "[FIT WINDOW", "leftmost_bar_index", "int"),
    ("Line Span Bars", "[FIT WINDOW", "line_span_bars", "int"),
    ("Baseline Coverage (n)", "[FIT WINDOW", "baseline_coverage_n", "int"),
    ("Baseline Coverage Rate", "[FIT WINDOW", "baseline_coverage_rate", "real"),
    ("Centroids Used", "[FIT WINDOW", "centroids_used", "int"),
    ("Crossings In Window (n)", "[FIT WINDOW", "crossings_in_window_n", "int"),
    ("First Crossing TS (UTC)", "[FIT WINDOW", "first_crossing_ts", "int"),
    ("Last Crossing TS (UTC)", "[FIT WINDOW", "last_crossing_ts", "int"),
    ("Live Close", None, "live_close", "real"),
    ("Baseline Value", "[PRICE CONTEXT", "baseline_value", "real"),
    ("UOEDT Value", "[PRICE CONTEXT", "uoedt_value", "real"),
    ("LOEDT Value", "[PRICE CONTEXT", "loedt_value", "real"),
    ("Distance To Baseline", "[PRICE CONTEXT", "dist_to_baseline", "real"),
    ("Distance To UOEDT", "[PRICE CONTEXT", "dist_to_uoedt", "real"),
    ("Distance To LOEDT", "[PRICE CONTEXT", "dist_to_loedt", "real"),
    ("Channel Position", "[PRICE CONTEXT", "channel_position", "real"),
    (("Window High", "Highest High"), None, "window_high", "real"),
    (("Window Low", "Lowest Low"), None, "window_low", "real"),
    ("Window Range", None, "window_range", "real"),
    ("Channel Width", "[CHANNEL GEOMETRY", "channel_width", "real"),
    ("Channel Asymmetry", "[CHANNEL GEOMETRY", "channel_asymmetry", "real"),
    ("Above UOEDT Count", "[CHANNEL GEOMETRY", "breach_above_n", "int"),
    ("Below LOEDT Count", "[CHANNEL GEOMETRY", "breach_below_n", "int"),
    ("Max Excursion Above", "[CHANNEL GEOMETRY", "max_excursion_above", "real"),
    ("Max Excursion Below", "[CHANNEL GEOMETRY", "max_excursion_below", "real"),
    ("Sample (n)", "[RESIDUAL DIAGNOSTICS; CROSSINGS", "resid_a_n", "int"),
    ("Mean Residual", "[RESIDUAL DIAGNOSTICS; CROSSINGS", "resid_a_mean", "real"),
    ("MAE", "[RESIDUAL DIAGNOSTICS; CROSSINGS", "resid_a_mae", "real"),
    ("Residual StdDev", "[RESIDUAL DIAGNOSTICS; CROSSINGS", "resid_a_sd", "real"),
    ("Max Abs Residual", "[RESIDUAL DIAGNOSTICS; CROSSINGS", "resid_a_max", "real"),
    ("Durbin-Watson", "[RESIDUAL DIAGNOSTICS; CROSSINGS", "resid_a_dw", "real"),
    ("Sample (n)", "[RESIDUAL DIAGNOSTICS; CLOSE", "resid_b_n", "int"),
    ("Mean Residual", "[RESIDUAL DIAGNOSTICS; CLOSE", "resid_b_mean", "real"),
    ("MAE", "[RESIDUAL DIAGNOSTICS; CLOSE", "resid_b_mae", "real"),
    ("Residual StdDev", "[RESIDUAL DIAGNOSTICS; CLOSE", "resid_b_sd", "real"),
    ("Max Abs Residual", "[RESIDUAL DIAGNOSTICS; CLOSE", "resid_b_max", "real"),
    ("Durbin-Watson", "[RESIDUAL DIAGNOSTICS; CLOSE", "resid_b_dw", "real"),
    # S&R Auto-Calibration Provenance (sr_levels source only)
    ("Fractals Sample (N)", "[SUPPORT-RESISTANCE", "sr_fractals_n", "int"),
    ("Q25 (25th percentile)", "[SUPPORT-RESISTANCE", "sr_q25", "real"),
    ("Q75 (75th percentile)", "[SUPPORT-RESISTANCE", "sr_q75", "real"),
    ("IQR", "[SUPPORT-RESISTANCE", "sr_iqr", "real"),
    ("Optimal Step", "[SUPPORT-RESISTANCE", "sr_optimal_step", "real"),
    ("Total Macro Clusters", "[SUPPORT-RESISTANCE", "sr_macro_clusters", "int"),
    ("Nearest Resistance", "[SUPPORT-RESISTANCE", "sr_nearest_resistance", "real"),
    ("Nearest Support", "[SUPPORT-RESISTANCE", "sr_nearest_support", "real"),
    ("Distance to Resistance", "[SUPPORT-RESISTANCE", "sr_dist_resistance_pts", "int"),
    ("Distance to Support", "[SUPPORT-RESISTANCE", "sr_dist_support_pts", "int"),
]

STAT_CONFIG_LABELS = {
    "Regression Centroids (Box B)",
    "Regression Centroids (Best WLS CFL)",
    "Excluded Recent Centroids (Box A)",
    "Excluded Centroids",
    "Time-Decay Lambda",
    "Visual EDT Window (Bars)",
    "Visual CFL/EDT Window (Bars)",
    "Timeframe (Sec)",
    "Fractal Bars",
    "Min Touches",
    "Require Both Sides",
    "Max Line Angle",
    "Tolerance Type",
    "Tolerance Percent",
    "Tolerance ATR Multiplier",
    "EDT Min Touches",
    "LOEDT Min Touches",
    "UOEDT Min Touches",
    "Extend To Current",
    "Projection Mode",
    "Frozen Anchor TS (Server)",
    "Frozen Slope (b)",
    "Frozen Anchor Price",
    "Frozen UOEDT Offset",
    "Frozen LOEDT Offset",
    "Calculation Mode",
    "Window Mode",
    "Min Touches Filter",
    "Max Window Bars",
}


def parse_statistic_files():
    stat_files = [
        ("Centriod_Best_Fit_A_XAUUSD_M5_Statistic.txt", "best_fit_a", "M5", 1789764900),
        ("Fractal_EDT_XAUUSD_M5_Statistic.txt", "fractal_edt", "M5", 1789764900),
        ("Non-Recent-B_XAUUSD_M15_Statistic.txt", "non_b", "M15", 1789764300),
        ("Resistance_Line_XAUUSD_M5_Statistic.txt", "resistance", "M5", 1789764900),
        ("Support_Line_XAUUSD_M5_Statistic.txt", "support", "M5", 1789764900),
        ("SR_Levels_XAUUSD_M15_Statistic.txt", "sr_levels", "M15", 1789764300),
    ]

    stat_rows = []
    for fname, source, tf, def_live_bar_ts in stat_files:
        p = os.path.join(DATA_DIR, fname)
        if not os.path.exists(p):
            continue
        with open(p, "r", encoding="utf-8", errors="replace") as f:
            text = f.read()

        seen = {}
        config = {}
        section = ""
        for line in text.splitlines():
            line = line.strip()
            if not line:
                continue
            if line.startswith("["):
                section = line
                continue
            if ":" not in line:
                continue
            label, _, value = line.partition(":")
            label, value = label.strip(), value.strip()
            seen[(section, label)] = value
            seen.setdefault((None, label), value)
            if label in STAT_CONFIG_LABELS:
                config[label] = value

        live_bar_ts = int(seen.get((None, "Live Bar TS (UTC)"), def_live_bar_ts))
        row = {
            "id": make_cuid("cs", f"stat_{source}_{tf}_{live_bar_ts}"),
            "terminal_id": "push_worker_v5",
            "symbol": "XAUUSD",
            "timeframe": tf,
            "source": source,
            "captured_at": 1789764900,
            "live_bar_ts": live_bar_ts,
            "cycle_id": 1,
            "createdAt": "2026-09-19 11:48:00",
        }

        canonical = json.dumps(config, sort_keys=True, separators=(",", ":"))
        row["config_hash"] = hashlib.sha256(canonical.encode("utf-8")).hexdigest()

        for label, sect_prefix, col, typ in STAT_FIELDS:
            raw = None
            for lb in (label if isinstance(label, tuple) else (label,)):
                if sect_prefix is None:
                    raw = seen.get((None, lb))
                else:
                    raw = next(
                        (
                            v
                            for (s, l), v in seen.items()
                            if l == lb and s is not None and s.startswith(sect_prefix)
                        ),
                        None,
                    )
                if raw is not None:
                    break

            if raw is not None and raw.strip() != "":
                raw_s = raw.strip()
                try:
                    if typ == "int":
                        row[col] = int(float(raw_s))
                    elif typ == "real":
                        row[col] = float(raw_s)
                    elif typ == "bool":
                        row[col] = (
                            True
                            if raw_s.lower() in ("true", "1")
                            else False
                            if raw_s.lower() in ("false", "0")
                            else None
                        )
                except ValueError:
                    row[col] = None
            else:
                if col not in row:
                    row[col] = None
        stat_rows.append(row)
    return stat_rows


def build_schema_dictionary():
    categories = [
        (1, 10, "Identity & Base Price Spine (OHLCV)", "MQL5 ohlcvexportlightweight_v2_29.mq5"),
        (11, 18, "Centroid Variant 1: best_fit_a", "MQL5 Centriod_Best_Fit_A + Python WLS"),
        (19, 26, "Centroid Variant 2: best_fit_b", "MQL5 Centriod_Best_Fit_B + Python WLS"),
        (27, 34, "Centroid Variant 3: cherry_a", "MQL5 CherryPickA + Python OLS"),
        (35, 42, "Centroid Variant 4: cherry_b", "MQL5 CherryPickB + Python OLS"),
        (43, 50, "Centroid Variant 5: most_recent", "MQL5 MostRecent + Python OLS"),
        (51, 58, "Centroid Variant 6: non_a", "MQL5 NonRecentA + Python OLS"),
        (59, 66, "Centroid Variant 7: non_b", "MQL5 NonRecentB + Python OLS"),
        (67, 71, "Fractal EDT & Single Best Lines", "MQL5 Fractal_EDT + Python fractal_lines.py"),
        (72, 79, "Support & Resistance Auto-Calibration (sr_1..sr_8)", "MQL5 SupportAndResistantAutoCalibration_v2_29.mq5"),
        (80, 82, "Z-Score Candle Volatility", "MQL5 zscoreohlccandleexport_v2_29.mq5 + Python"),
        (83, 93, "ZigZag Market Structure & SMC Metrics", "MQL5 ZigZagExportv43_v2_29.mq5 + Python zigzag.py"),
        (94, 96, "Provenance Tracking", "Contabo VPS Collection Pipeline"),
        (97, 98, "Record Timestamps", "Prisma ORM System Timestamps"),
    ]

    field_meta = {
        "id": ("String", "TEXT", "REQUIRED", "Primary key: cuid unique identifier for each bar record"),
        "terminal_id": ("String", "TEXT", "REQUIRED", "Sender ID; 'push_worker_v5' from the v6 push worker"),
        "timestamp": ("Int", "INTEGER", "REQUIRED", "Adjusted bar timestamp in UTC unix seconds, aligned to timeframe grid (300s/900s)"),
        "symbol": ("String", "TEXT", "REQUIRED", "Asset symbol; locked to 'XAUUSD' (Gold / US Dollar)"),
        "timeframe": ("String", "TEXT", "REQUIRED", "Bar timeframe: 'M5' (Micro Execution) or 'M15' (Macro Structure)"),
        "open": ("Float", "DOUBLE PRECISION", "REQUIRED", "Bar Open price in USD per troy ounce"),
        "high": ("Float", "DOUBLE PRECISION", "REQUIRED", "Bar High price in USD per troy ounce"),
        "low": ("Float", "DOUBLE PRECISION", "REQUIRED", "Bar Low price in USD per troy ounce"),
        "close": ("Float", "DOUBLE PRECISION", "REQUIRED", "Bar Close price in USD per troy ounce"),
        "volume": ("Int", "INTEGER", "REQUIRED", "Tick volume recorded for this bar"),
        # best_fit_a
        "best_fit_a_horiz_high_map": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Highest upper fractal peak marker in observation cycle (symbol 108)"),
        "best_fit_a_horiz_low_map": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Lowest lower fractal valley marker in observation cycle (symbol 108)"),
        "best_fit_a_ssa": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Singular Spectrum Analysis (SSA Window=30, Rank=6) denoised trendline"),
        "best_fit_a_ema_ssa": ("Float?", "DOUBLE PRECISION", "NULLABLE", "EMA of SSA (Signal=3) trigger line for crossover detection"),
        "best_fit_a_crossing": ("Int?", "INTEGER", "NULLABLE", "SSA x EMA_SSA crossover event flag: 1 = crossover on this bar, 0 = no crossover"),
        "best_fit_a_base_fl": ("Float?", "DOUBLE PRECISION", "NULLABLE", "WLS Centroid Regression Flip Line (Central Baseline, Lambda=0.000, Exclude=0)"),
        "best_fit_a_uoedt": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Upper Outermost Equidistant Trendline (Overbought / Resistance corridor)"),
        "best_fit_a_loedt": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Lower Outermost Equidistant Trendline (Oversold / Support corridor)"),
        # best_fit_b
        "best_fit_b_horiz_high_map": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Highest upper fractal marker for secondary Best-Fit B preset"),
        "best_fit_b_horiz_low_map": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Lowest lower fractal marker for secondary Best-Fit B preset"),
        "best_fit_b_ssa": ("Float?", "DOUBLE PRECISION", "NULLABLE", "SSA trendline for secondary Best-Fit B preset"),
        "best_fit_b_ema_ssa": ("Float?", "DOUBLE PRECISION", "NULLABLE", "EMA of SSA line for secondary Best-Fit B preset"),
        "best_fit_b_crossing": ("Int?", "INTEGER", "NULLABLE", "Crossover trigger flag for Best-Fit B preset"),
        "best_fit_b_base_fl": ("Float?", "DOUBLE PRECISION", "NULLABLE", "WLS Baseline excluding 3 most recent centroids to avoid short-term lag"),
        "best_fit_b_uoedt": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Upper Outermost EDT channel boundary for Best-Fit B preset"),
        "best_fit_b_loedt": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Lower Outermost EDT channel boundary for Best-Fit B preset"),
        # cherry_a
        "cherry_a_horiz_high_map": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Cherry-Pick Group A upper fractal level marker"),
        "cherry_a_horiz_low_map": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Cherry-Pick Group A lower fractal level marker"),
        "cherry_a_ssa": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Cherry-Pick Group A SSA trendline"),
        "cherry_a_ema_ssa": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Cherry-Pick Group A EMA SSA trigger line"),
        "cherry_a_crossing": ("Int?", "INTEGER", "NULLABLE", "Cherry-Pick Group A crossover flag"),
        "cherry_a_base_fl": ("Float?", "DOUBLE PRECISION", "NULLABLE", "OLS Regression Baseline fitted to Cherry-Pick Group A centroids"),
        "cherry_a_uoedt": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Cherry-Pick Group A Upper EDT channel"),
        "cherry_a_loedt": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Cherry-Pick Group A Lower EDT channel"),
        # cherry_b
        "cherry_b_horiz_high_map": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Cherry-Pick Group B upper fractal level marker"),
        "cherry_b_horiz_low_map": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Cherry-Pick Group B lower fractal level marker"),
        "cherry_b_ssa": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Cherry-Pick Group B SSA trendline"),
        "cherry_b_ema_ssa": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Cherry-Pick Group B EMA SSA trigger line"),
        "cherry_b_crossing": ("Int?", "INTEGER", "NULLABLE", "Cherry-Pick Group B crossover flag"),
        "cherry_b_base_fl": ("Float?", "DOUBLE PRECISION", "NULLABLE", "OLS Regression Baseline fitted to Cherry-Pick Group B centroids"),
        "cherry_b_uoedt": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Cherry-Pick Group B Upper EDT channel"),
        "cherry_b_loedt": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Cherry-Pick Group B Lower EDT channel"),
        # most_recent
        "most_recent_horiz_high_map": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Most-Recent variant upper fractal level marker"),
        "most_recent_horiz_low_map": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Most-Recent variant lower fractal level marker"),
        "most_recent_ssa": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Most-Recent variant SSA trendline"),
        "most_recent_ema_ssa": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Most-Recent variant EMA SSA trigger line"),
        "most_recent_crossing": ("Int?", "INTEGER", "NULLABLE", "Most-Recent variant crossover flag"),
        "most_recent_base_fl": ("Float?", "DOUBLE PRECISION", "NULLABLE", "OLS Regression Baseline fitted to N most recent centroids"),
        "most_recent_uoedt": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Most-Recent variant Upper EDT channel"),
        "most_recent_loedt": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Most-Recent variant Lower EDT channel"),
        # non_a
        "non_a_horiz_high_map": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Non-Recent Group A upper fractal level marker"),
        "non_a_horiz_low_map": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Non-Recent Group A lower fractal level marker"),
        "non_a_ssa": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Non-Recent Group A SSA trendline"),
        "non_a_ema_ssa": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Non-Recent Group A EMA SSA trigger line"),
        "non_a_crossing": ("Int?", "INTEGER", "NULLABLE", "Non-Recent Group A crossover flag"),
        "non_a_base_fl": ("Float?", "DOUBLE PRECISION", "NULLABLE", "OLS Regression Baseline excluding recent centroids Group A"),
        "non_a_uoedt": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Non-Recent Group A Upper EDT channel"),
        "non_a_loedt": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Non-Recent Group A Lower EDT channel"),
        # non_b
        "non_b_horiz_high_map": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Non-Recent Group B upper fractal level marker"),
        "non_b_horiz_low_map": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Non-Recent Group B lower fractal level marker"),
        "non_b_ssa": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Non-Recent Group B SSA trendline"),
        "non_b_ema_ssa": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Non-Recent Group B EMA SSA trigger line"),
        "non_b_crossing": ("Int?", "INTEGER", "NULLABLE", "Non-Recent Group B crossover flag"),
        "non_b_base_fl": ("Float?", "DOUBLE PRECISION", "NULLABLE", "OLS Regression Baseline excluding recent centroids Group B"),
        "non_b_uoedt": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Non-Recent Group B Upper EDT channel"),
        "non_b_loedt": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Non-Recent Group B Lower EDT channel"),
        # fractal & lines
        "fractal_best_fl": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Fractal Best Flip Line: axis line connecting highest touch upper/lower fractals"),
        "fractal_uoedt": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Upper Outermost EDT channel parallel to Fractal Best FL"),
        "fractal_loedt": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Lower Outermost EDT channel parallel to Fractal Best FL"),
        "best_resistance": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Single Best Resistance: dominant upper resistance line with max touch score"),
        "best_support": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Single Best Support: dominant lower support line with max touch score"),
        # SR Levels
        "sr_1": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Auto-calibrated Support 1: closest support level BELOW bar close"),
        "sr_2": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Auto-calibrated Support 2: 2nd closest support level below bar close"),
        "sr_3": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Auto-calibrated Support 3: 3rd closest support level below bar close"),
        "sr_4": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Auto-calibrated Support 4: 4th closest support level below bar close"),
        "sr_5": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Auto-calibrated Resistance 1: closest resistance level ABOVE bar close"),
        "sr_6": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Auto-calibrated Resistance 2: 2nd closest resistance level above bar close"),
        "sr_7": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Auto-calibrated Resistance 3: 3rd closest resistance level above bar close"),
        "sr_8": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Auto-calibrated Resistance 4: 4th closest resistance level above bar close"),
        # Z-Score
        "body_direction": ("Int?", "INTEGER", "NULLABLE", "Candle direction: +1 (Bullish/Green), -1 (Bearish/Red), 0 (Doji)"),
        "body_size": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Normalized candle body size |Z-score| relative to 432-bar rolling window"),
        "body_classification": ("Int?", "INTEGER", "NULLABLE", "Volatility classification (0: UP_NORMAL, 1: UP_LARGE, 2: UP_EXTREME, 3: DN_NORMAL, 4: DN_LARGE, 5: DN_EXTREME)"),
        # ZigZag
        "zigzag_point_type": ("String?", "TEXT", "NULLABLE", "ZigZag swing pivot type: 'Peak' (Upper swing) or 'Bottom' (Lower swing); NULL on non-pivot bars"),
        "zigzag_current_point": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Price level at confirmed swing pivot"),
        "zigzag_price_change": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Absolute price delta from previous swing (P_curr - P_prev)"),
        "zigzag_pct_change": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Percentage price change from previous swing"),
        "zigzag_pct_change_class": ("Int?", "INTEGER", "NULLABLE", "Z-score classification of percentage price change (0-2 up, 3-5 down)"),
        "zigzag_bars": ("Int?", "INTEGER", "NULLABLE", "Number of bars elapsed during this wave segment (Wave duration)"),
        "zigzag_bars_class": ("Int?", "INTEGER", "NULLABLE", "Z-score classification of wave bar duration (0-5)"),
        "zigzag_price_per_bar": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Price velocity per bar (|Price delta| / Bars)"),
        "zigzag_price_per_bar_class": ("Int?", "INTEGER", "NULLABLE", "Z-score classification of price velocity (0-5)"),
        "zigzag_slope": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Geometric wave slope in degrees (arctan(PricePerBar) * 180 / pi)"),
        "zigzag_category": ("String?", "TEXT", "NULLABLE", "Smart Money Concept market structure: HH, HL, LH, LL, EQH, EQL"),
        # Provenance
        "cycle_id": ("Int", "INTEGER", "REQUIRED", "Ingestion cycle tracking ID from collection_cycles table"),
        "collected_at": ("Int", "INTEGER", "REQUIRED", "Unix epoch timestamp (UTC seconds) when raw MQL5 export files were ingested"),
        "calculated_at": ("Int?", "INTEGER", "NULLABLE", "Unix epoch timestamp (UTC seconds) when Python derived metrics completed"),
        "createdAt": ("DateTime", "TIMESTAMPTZ", "REQUIRED", "System record creation timestamp"),
        "updatedAt": ("DateTime", "TIMESTAMPTZ", "REQUIRED", "System record last update timestamp"),
    }

    dict_rows = []
    for idx, col in enumerate(MARKET_DATA_V6_COLUMNS, start=1):
        cat_name = "Other"
        cat_origin = "System"
        for start_c, end_c, name, origin in categories:
            if start_c <= idx <= end_c:
                cat_name = name
                cat_origin = origin
                break
        meta = field_meta.get(col, ("Unknown", "TEXT", "NULLABLE", "Indicator field"))
        dict_rows.append({
            "col_index": idx,
            "column_name": col,
            "prisma_type": meta[0],
            "postgres_type": meta[1],
            "nullability": meta[2],
            "category": cat_name,
            "origin_source": cat_origin,
            "description": meta[3],
        })
    return dict_rows


def build_indicator_statistic_dictionary():
    categories = [
        (1, 8, "Identity & Ingestion Timing", "Pipeline Ingestion Architecture"),
        (9, 14, "Resolved Geometric Line", "MQL5 BestFL & Trendline Indicators"),
        (15, 19, "Fitting Window Definition", "MQL5 Calculation Window"),
        (20, 25, "Model A Crossings Fit Quality", "MQL5 SSA Crossings Regression"),
        (26, 31, "Model B Close Price Fit Quality", "MQL5 Close Price Regression"),
        (32, 36, "EDT Channel Geometry & Containment", "MQL5 Equidistant Channel Envelope"),
        (37, 47, "Extended Fit Window Provenance", "MQL5 Extended Fit Window (2026-09-20)"),
        (48, 58, "Extended Price Context (Live Bar)", "MQL5 Extended Price Context (2026-09-20)"),
        (59, 64, "Extended Channel Geometry & Excursions", "MQL5 Channel Geometry (2026-09-20)"),
        (65, 70, "Residual Diagnostics: Model A Crossings", "MQL5 Crossings Diagnostics (2026-09-20)"),
        (71, 76, "Residual Diagnostics: Model B Close Price", "MQL5 Close Price Diagnostics (2026-09-20)"),
        (77, 86, "Support & Resistance Auto-Calibration", "MQL5 14th Indicator: Freedman-Diaconis Clustering"),
        (87, 88, "Configuration Hash & Record Timestamps", "Prisma ORM & Security Governance"),
    ]

    stat_meta = {
        "id": ("String", "TEXT", "REQUIRED", "Primary key: cuid unique identifier for each statistic snapshot"),
        "terminal_id": ("String", "TEXT", "REQUIRED", "Originating terminal / push worker identifier ('push_worker_v5')"),
        "symbol": ("String", "TEXT", "REQUIRED", "Financial instrument symbol ('XAUUSD')"),
        "timeframe": ("String", "TEXT", "REQUIRED", "Chart timeframe of indicator calculation ('M5' or 'M15')"),
        "source": ("String", "TEXT", "REQUIRED", "Indicator source variant (best_fit_a, non_b, fractal_edt, resistance, support, sr_levels, etc.)"),
        "captured_at": ("Int", "INTEGER", "REQUIRED", "Unix UTC timestamp of collector's 5-minute cycle slot; append-only key half"),
        "live_bar_ts": ("Int", "INTEGER", "REQUIRED", "Unix UTC timestamp of newest bar in export; joins to MarketDataV6.timestamp"),
        "cycle_id": ("Int", "INTEGER", "REQUIRED", "Collector cycle identification number"),
        "solution_found": ("Boolean?", "BOOLEAN", "NULLABLE", "Boolean flag indicating if geometric line solver resolved a valid fit"),
        "raw_slope": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Raw mathematical slope (b) of baseline in price units per bar"),
        "anchored_y_int": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Y-intercept price value anchored at the live bar"),
        "regression_angle": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Regression angle in degrees (-90 to +90) relative to chart aspect ratio"),
        "line_origin_ts": ("Int?", "INTEGER", "NULLABLE", "Unix UTC timestamp of earliest anchor / origin point of resolved line"),
        "touches": ("Int?", "INTEGER", "NULLABLE", "Total count of fractal touches confirming the resolved trendline"),
        "window_start_ts": ("Int?", "INTEGER", "NULLABLE", "Unix UTC timestamp of oldest bar in evaluated observation window"),
        "window_end_ts": ("Int?", "INTEGER", "NULLABLE", "Unix UTC timestamp of newest bar in evaluated observation window"),
        "window_bars": ("Int?", "INTEGER", "NULLABLE", "Observation window bar count (Box B Bars) used for centroid evaluation"),
        "math_lookback": ("Int?", "INTEGER", "NULLABLE", "Search window bar count searched by indicator (Math Search Window)"),
        "crossings_n": ("Int?", "INTEGER", "NULLABLE", "Total SSA x EMA-SSA crossover points identified in search window"),
        "model_a_n": ("Int?", "INTEGER", "NULLABLE", "Model A sample size: number of centroid crossing points fitted"),
        "model_a_r2": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Model A Coefficient of Determination (R²) against centroid crossing points"),
        "model_a_mse": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Model A Mean Squared Error against centroid crossing points"),
        "model_a_var_ratio": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Model A Variance Ratio (residual variance / total variance)"),
        "model_a_skew": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Model A Skewness of residuals against centroid points"),
        "model_a_kurt": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Model A Kurtosis of residuals against centroid points"),
        "model_b_n": ("Int?", "INTEGER", "NULLABLE", "Model B sample size: number of candlestick close prices evaluated"),
        "model_b_r2": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Model B Coefficient of Determination (R²) against close prices"),
        "model_b_mse": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Model B Mean Squared Error against close prices"),
        "model_b_var_ratio": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Model B Variance Ratio against close prices"),
        "model_b_skew": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Model B Skewness of residuals against close prices"),
        "model_b_kurt": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Model B Kurtosis of residuals against close prices"),
        "uoedt_offset": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Distance in USD from baseline to Upper Outermost EDT channel boundary"),
        "loedt_offset": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Distance in USD from baseline to Lower Outermost EDT channel boundary (negative)"),
        "containment_n": ("Int?", "INTEGER", "NULLABLE", "Containment sample size: number of bars evaluated for channel containment"),
        "containment_count": ("Int?", "INTEGER", "NULLABLE", "Number of candlestick closes fully contained within EDT channel"),
        "containment_rate": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Percentage of bars contained within EDT channel (100 * count / n)"),
        "window_span_bars": ("Int?", "INTEGER", "NULLABLE", "Bars from leftmost evaluated bar to live bar inclusive"),
        "visual_window_bars": ("Int?", "INTEGER", "NULLABLE", "Bars channel is drawn over on chart (can exceed evaluated window)"),
        "bars_available": ("Int?", "INTEGER", "NULLABLE", "Total bars indicator actually had in terminal buffer"),
        "leftmost_bar_index": ("Int?", "INTEGER", "NULLABLE", "Index of oldest evaluated bar for reproducing fit"),
        "line_span_bars": ("Int?", "INTEGER", "NULLABLE", "Bars between oldest and newest resolved baseline value"),
        "baseline_coverage_n": ("Int?", "INTEGER", "NULLABLE", "Bars in span carrying a resolved baseline value"),
        "baseline_coverage_rate": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Baseline coverage percentage (100 * coverage_n / line_span_bars)"),
        "centroids_used": ("Int?", "INTEGER", "NULLABLE", "Centroid clusters actually used in winning regression fit"),
        "crossings_in_window_n": ("Int?", "INTEGER", "NULLABLE", "SSA crossover points located inside drawn visual span"),
        "first_crossing_ts": ("Int?", "INTEGER", "NULLABLE", "Unix UTC timestamp of oldest SSA crossing in drawn span"),
        "last_crossing_ts": ("Int?", "INTEGER", "NULLABLE", "Unix UTC timestamp of newest SSA crossing in drawn span"),
        "live_close": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Close price of newest bar (self-contained snapshot)"),
        "baseline_value": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Fitted baseline value AT the live bar"),
        "uoedt_value": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Upper EDT channel boundary price AT the live bar"),
        "loedt_value": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Lower EDT channel boundary price AT the live bar"),
        "dist_to_baseline": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Distance from live close to baseline (live_close - baseline_value)"),
        "dist_to_uoedt": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Distance from live close to UOEDT (uoedt_value - live_close)"),
        "dist_to_loedt": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Distance from live close to LOEDT (live_close - loedt_value)"),
        "channel_position": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Relative price position in channel (0.0 at LOEDT, 1.0 at UOEDT)"),
        "window_high": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Highest high price across resolved span"),
        "window_low": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Lowest low price across resolved span"),
        "window_range": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Price range across window (window_high - window_low)"),
        "channel_width": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Channel width in USD (uoedt_value - loedt_value)"),
        "channel_asymmetry": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Channel asymmetry ratio: (uoedt_offset + loedt_offset) / channel_width"),
        "breach_above_n": ("Int?", "INTEGER", "NULLABLE", "Count of bar closes breaching above upper EDT boundary"),
        "breach_below_n": ("Int?", "INTEGER", "NULLABLE", "Count of bar closes breaching below lower EDT boundary"),
        "max_excursion_above": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Maximum distance of close above UOEDT during breach"),
        "max_excursion_below": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Maximum distance of close below LOEDT during breach"),
        "resid_a_n": ("Int?", "INTEGER", "NULLABLE", "Sample size for crossings residual diagnostics over drawn span"),
        "resid_a_mean": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Mean residual (bias) of baseline against crossings"),
        "resid_a_mae": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Mean Absolute Error of baseline against crossings"),
        "resid_a_sd": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Residual Standard Deviation around baseline for crossings"),
        "resid_a_max": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Maximum absolute residual miss against crossings"),
        "resid_a_dw": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Durbin-Watson serial correlation statistic for crossings residuals"),
        "resid_b_n": ("Int?", "INTEGER", "NULLABLE", "Sample size for close price residual diagnostics"),
        "resid_b_mean": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Mean residual (bias) of baseline against close prices"),
        "resid_b_mae": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Mean Absolute Error of baseline against close prices"),
        "resid_b_sd": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Residual Standard Deviation around baseline for close prices"),
        "resid_b_max": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Maximum absolute residual miss against close prices"),
        "resid_b_dw": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Durbin-Watson serial correlation statistic for close price residuals"),
        "sr_fractals_n": ("Int?", "INTEGER", "NULLABLE", "Number of fractals used in Freedman-Diaconis calibration sample"),
        "sr_q25": ("Float?", "DOUBLE PRECISION", "NULLABLE", "25th percentile price level of fractal sample (Q25)"),
        "sr_q75": ("Float?", "DOUBLE PRECISION", "NULLABLE", "75th percentile price level of fractal sample (Q75)"),
        "sr_iqr": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Interquartile Range of fractals (Q75 - Q25)"),
        "sr_optimal_step": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Auto-calibrated optimal cluster bucket step size (2 * IQR * N^(-1/3))"),
        "sr_macro_clusters": ("Int?", "INTEGER", "NULLABLE", "Total macro clusters identified before 8-slot truncation"),
        "sr_nearest_resistance": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Auto-calibrated resistance level closest above live close"),
        "sr_nearest_support": ("Float?", "DOUBLE PRECISION", "NULLABLE", "Auto-calibrated support level closest below live close"),
        "sr_dist_resistance_pts": ("Int?", "INTEGER", "NULLABLE", "Distance to nearest resistance in points (0.01 USD per point)"),
        "sr_dist_support_pts": ("Int?", "INTEGER", "NULLABLE", "Distance to nearest support in points (0.01 USD per point)"),
        "config_hash": ("String", "TEXT", "REQUIRED", "SHA256 hash of normalized indicator configuration parameters"),
        "createdAt": ("DateTime", "TIMESTAMPTZ", "REQUIRED", "Timestamp when record was created in database"),
    }

    dict_rows = []
    for idx, col in enumerate(INDICATOR_STATISTIC_COLUMNS, start=1):
        cat_name = "Other"
        cat_origin = "System"
        for start_c, end_c, name, origin in categories:
            if start_c <= idx <= end_c:
                cat_name = name
                cat_origin = origin
                break
        meta = stat_meta.get(col, ("Unknown", "TEXT", "NULLABLE", "Statistic field"))
        dict_rows.append({
            "col_index": idx,
            "column_name": col,
            "prisma_type": meta[0],
            "postgres_type": meta[1],
            "nullability": meta[2],
            "category": cat_name,
            "origin_source": cat_origin,
            "description": meta[3],
        })
    return dict_rows


def create_excel():
    print("Loading data from engine-1-5-new...", flush=True)
    records = load_all_market_data()
    print(f"Loaded {len(records)} total records.", flush=True)

    print("Parsing statistic files...", flush=True)
    stat_rows = parse_statistic_files()
    print(f"Parsed {len(stat_rows)} statistic rows.", flush=True)

    print("Building schema dictionaries...", flush=True)
    schema_dict_market = build_schema_dictionary()
    schema_dict_stat = build_indicator_statistic_dictionary()

    print("Creating Excel workbook...", flush=True)
    wb = openpyxl.Workbook()

    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_data = Font(name="Calibri", size=10)
    font_bold = Font(name="Calibri", size=10, bold=True)

    fill_navy = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    fill_blue_header = PatternFill(start_color="2C5282", end_color="2C5282", fill_type="solid")
    fill_teal_header = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
    fill_subtle_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    border_thin = Side(style="thin", color="CBD5E1")
    cell_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    align_header = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def populate_market_data_sheet(ws, title, rows_dict):
        print(f"Populating sheet: {title} ({len(rows_dict)} rows)...", flush=True)
        ws.title = title
        ws.views.sheetView[0].showGridLines = True

        # Header Row
        ws.append(MARKET_DATA_V6_COLUMNS)
        ws.row_dimensions[1].height = 28

        for col_num in range(1, len(MARKET_DATA_V6_COLUMNS) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = font_header
            cell.fill = fill_navy
            cell.alignment = align_header
            cell.border = cell_border

        # Sort keys by timeframe, then timestamp
        sorted_keys = sorted(rows_dict.keys(), key=lambda k: (k[0], k[1]))

        # Append all rows quickly
        for r_idx, key in enumerate(sorted_keys, start=2):
            row_data = rows_dict[key]
            row_vals = [row_data.get(col, None) for col in MARKET_DATA_V6_COLUMNS]
            ws.append(row_vals)

        # Freeze Panes: freeze row 1 and first 5 columns (id, terminal_id, timestamp, symbol, timeframe)
        ws.freeze_panes = "F2"

        # Set column widths
        for col_idx, col_name in enumerate(MARKET_DATA_V6_COLUMNS, start=1):
            col_letter = get_column_letter(col_idx)
            w = max(len(col_name) + 3, 12)
            if col_name == "id":
                w = 27
            elif col_name in ("createdAt", "updatedAt"):
                w = 20
            elif col_name == "terminal_id":
                w = 16
            ws.column_dimensions[col_letter].width = min(w, 32)

    # 1. Sheet 1: market_data_v6 (All 6,000 rows)
    ws1 = wb.active
    populate_market_data_sheet(ws1, "market_data_v6", records)

    # 2. Sheet 2: market_data_v6_M5 (3,000 rows)
    m5_records = {k: v for k, v in records.items() if k[0] == "M5"}
    ws2 = wb.create_sheet("market_data_v6_M5")
    populate_market_data_sheet(ws2, "market_data_v6_M5", m5_records)

    # 3. Sheet 3: market_data_v6_M15 (3,000 rows)
    m15_records = {k: v for k, v in records.items() if k[0] == "M15"}
    ws3 = wb.create_sheet("market_data_v6_M15")
    populate_market_data_sheet(ws3, "market_data_v6_M15", m15_records)

    # 4. Sheet 4: indicator_statistics (6 rows, 88 columns)
    print("Populating sheet: indicator_statistics (88 columns)...", flush=True)
    ws4 = wb.create_sheet("indicator_statistics")
    ws4.views.sheetView[0].showGridLines = True
    ws4.append(INDICATOR_STATISTIC_COLUMNS)
    ws4.row_dimensions[1].height = 28

    for c_idx in range(1, len(INDICATOR_STATISTIC_COLUMNS) + 1):
        cell = ws4.cell(row=1, column=c_idx)
        cell.font = font_header
        cell.fill = fill_blue_header
        cell.alignment = align_header
        cell.border = cell_border

    INT_STAT_COLUMNS = {
        "cycle_id",
        "captured_at",
        "live_bar_ts",
        "line_origin_ts",
        "touches",
        "window_start_ts",
        "window_end_ts",
        "window_bars",
        "math_lookback",
        "crossings_n",
        "model_a_n",
        "model_b_n",
        "containment_n",
        "containment_count",
        "window_span_bars",
        "visual_window_bars",
        "bars_available",
        "leftmost_bar_index",
        "line_span_bars",
        "baseline_coverage_n",
        "centroids_used",
        "crossings_in_window_n",
        "first_crossing_ts",
        "last_crossing_ts",
        "breach_above_n",
        "breach_below_n",
        "resid_a_n",
        "resid_b_n",
        "sr_fractals_n",
        "sr_macro_clusters",
        "sr_dist_resistance_pts",
        "sr_dist_support_pts",
    }
    FOUR_DEC_COLUMNS = {
        "raw_slope",
        "model_a_r2",
        "model_b_r2",
        "channel_position",
        "channel_asymmetry",
        "resid_a_dw",
        "resid_b_dw",
        "model_a_var_ratio",
        "model_b_var_ratio",
    }

    for r_idx, srow in enumerate(stat_rows, start=2):
        row_vals = [srow.get(c, None) for c in INDICATOR_STATISTIC_COLUMNS]
        ws4.append(row_vals)
        ws4.row_dimensions[r_idx].height = 20
        is_zebra = r_idx % 2 == 0
        row_fill = fill_subtle_zebra if is_zebra else fill_white

        for c_idx, val in enumerate(row_vals, start=1):
            col_name = INDICATOR_STATISTIC_COLUMNS[c_idx - 1]
            cell = ws4.cell(row=r_idx, column=c_idx)
            cell.font = font_data
            cell.fill = row_fill
            cell.border = cell_border
            if val is None:
                cell.alignment = align_center
            elif isinstance(val, bool):
                cell.alignment = align_center
            elif isinstance(val, (int, float)):
                cell.alignment = align_right
                if col_name in INT_STAT_COLUMNS or isinstance(val, int):
                    cell.number_format = "0"
                elif col_name in FOUR_DEC_COLUMNS:
                    cell.number_format = "0.0000"
                else:
                    cell.number_format = "0.00"
            else:
                cell.alignment = align_left

    ws4.freeze_panes = "F2"
    for col_idx, col_name in enumerate(INDICATOR_STATISTIC_COLUMNS, start=1):
        col_letter = get_column_letter(col_idx)
        w = max(len(col_name) + 3, 12)
        if col_name in ("id", "config_hash"):
            w = 28
        elif col_name in ("captured_at", "live_bar_ts", "createdAt"):
            w = 20
        ws4.column_dimensions[col_letter].width = min(w, 32)

    # Helper for rendering dictionary sheets
    def populate_dictionary_sheet(ws, title, header_fill, dict_rows):
        print(f"Populating sheet: {title} ({len(dict_rows)} rows)...", flush=True)
        ws.title = title
        ws.views.sheetView[0].showGridLines = True
        dict_headers = [
            "Col #",
            "Column Name",
            "Prisma Type",
            "PostgreSQL Type",
            "Nullability",
            "Category",
            "Source Indicator / Component",
            "Description & Mathematical Role",
        ]
        ws.append(dict_headers)
        ws.row_dimensions[1].height = 28

        for c_idx in range(1, len(dict_headers) + 1):
            cell = ws.cell(row=1, column=c_idx)
            cell.font = font_header
            cell.fill = header_fill
            cell.alignment = align_header
            cell.border = cell_border

        fill_req = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
        fill_null = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")

        for r_idx, drow in enumerate(dict_rows, start=2):
            row_vals = [
                drow["col_index"],
                drow["column_name"],
                drow["prisma_type"],
                drow["postgres_type"],
                drow["nullability"],
                drow["category"],
                drow["origin_source"],
                drow["description"],
            ]
            ws.append(row_vals)
            ws.row_dimensions[r_idx].height = 22
            is_zebra = r_idx % 2 == 0
            base_fill = fill_subtle_zebra if is_zebra else fill_white

            for c_idx, val in enumerate(row_vals, start=1):
                cell = ws.cell(row=r_idx, column=c_idx)
                cell.font = font_bold if c_idx == 2 else font_data
                cell.border = cell_border
                cell.fill = base_fill

                if c_idx == 1:
                    cell.alignment = align_center
                elif c_idx in (2, 3, 4):
                    cell.alignment = align_left
                elif c_idx == 5:
                    cell.alignment = align_center
                    cell.fill = fill_req if val == "REQUIRED" else fill_null
                    cell.font = Font(
                        name="Calibri",
                        size=10,
                        bold=True,
                        color="B45309" if val == "REQUIRED" else "1D4ED8",
                    )
                elif c_idx in (6, 7):
                    cell.alignment = align_left
                else:
                    cell.alignment = align_left

        ws.freeze_panes = "C2"
        ws.column_dimensions["A"].width = 8
        ws.column_dimensions["B"].width = 30
        ws.column_dimensions["C"].width = 16
        ws.column_dimensions["D"].width = 20
        ws.column_dimensions["E"].width = 15
        ws.column_dimensions["F"].width = 36
        ws.column_dimensions["G"].width = 44
        ws.column_dimensions["H"].width = 72

    # 5. Sheet 5: Schema_Dictionary (MarketDataV6 - 98 fields)
    ws5 = wb.create_sheet("Schema_Dictionary")
    populate_dictionary_sheet(ws5, "Schema_Dictionary", fill_navy, schema_dict_market)

    # 6. Sheet 6: Schema_Dict_IndicatorStat (IndicatorStatistic - 88 fields)
    ws6 = wb.create_sheet("Schema_Dict_IndicatorStat")
    populate_dictionary_sheet(ws6, "Schema_Dict_IndicatorStat", fill_teal_header, schema_dict_stat)

    # Save files
    saved_paths = []
    print(f"Saving to {OUTPUT_FILE_1}...", flush=True)
    try:
        wb.save(OUTPUT_FILE_1)
        saved_paths.append(OUTPUT_FILE_1)
        print(f"Successfully saved to {OUTPUT_FILE_1}", flush=True)
    except PermissionError:
        alt_output = os.path.join(DATA_DIR, "market_data_v6_replicated_updated.xlsx")
        print(f"NOTICE: {OUTPUT_FILE_1} is currently open in Microsoft Excel.", flush=True)
        print(f"Saving to alternate file: {alt_output}...", flush=True)
        wb.save(alt_output)
        saved_paths.append(alt_output)
        print(f"Successfully saved to {alt_output}", flush=True)

    print(f"Saving copy to {OUTPUT_FILE_2}...", flush=True)
    try:
        wb.save(OUTPUT_FILE_2)
        saved_paths.append(OUTPUT_FILE_2)
        print(f"Successfully saved to {OUTPUT_FILE_2}", flush=True)
    except PermissionError:
        print(f"NOTICE: {OUTPUT_FILE_2} is also locked by Excel.", flush=True)

    print(f"SUCCESS: Excel files created successfully with all 6 sheets! Saved: {saved_paths}", flush=True)


if __name__ == "__main__":
    create_excel()
